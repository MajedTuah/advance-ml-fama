"""Builds the Day 1 (4-hour) deck in the Session 05 visual style.

Usage:  python build_deck.py
PDF:    DECK_MONO_FONT=Menlo DECK_OUT=/tmp/day1.pptx python build_deck.py
        soffice --headless --convert-to pdf /tmp/day1.pptx
Needs:  python-pptx, numpy, scipy
"""
import os
from pathlib import Path

import numpy as np
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from scipy.stats import ks_2samp

HERE = Path(__file__).parent
ASSETS = HERE / "assets"
OUT = Path(os.environ.get("DECK_OUT", HERE / "Day 1 - Advanced ML in Production (4 Hours).pptx"))

# ---------------------------------------------------------------- theme (Session 05)
BG = "0F1A2E"
CARD = "172744"
CARD2 = "1E2F50"
PILL = "2A355C"
CODE_BG = "0B1424"
LINE = "2A3A5C"
ACCENT = "547AFF"
CYAN = "4FC3F7"
GREEN = "4CAF50"
ORANGE = "FF9800"
RED = "EF5350"
INDIGO = "7986CB"
WHITE = "FFFFFF"
MUTED = "A0AEC4"
DIM = "6B8EA8"
CODE_FG = "D6E2F5"
CODE_CMT = "7FB77E"

BODY = "Calibri"
# Consolas ships with Office; machines without it (e.g. LibreOffice on macOS) export the PDF with Menlo.
MONO = os.environ.get("DECK_MONO_FONT", "Consolas")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
_slide_no = [0]


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- primitives
def rect(slide, x, y, w, h, fill=CARD, radius=0.12, line=None, name=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if name:
        s.name = name
    s.text_frame.text = ""
    return s


def oval(slide, x, y, d, fill, text="", color=WHITE, size=14):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = text
    _style(r, size=size, color=color, bold=True)
    return s


def _style(run, size=14, color=WHITE, bold=False, italic=False, font=BODY):
    f = run.font
    f.size = Pt(size)
    f.color.rgb = rgb(color)
    f.bold = bold
    f.italic = italic
    f.name = font


def text(slide, x, y, w, h, paras, size=14, color=MUTED, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4, margin=0.0, line_spacing=None):
    """paras: str | list of paragraphs. A paragraph is a str, or a list of runs.
    A run is a str or (text, {size,color,bold,italic,font}). A paragraph may also be
    a dict {"runs": [...], "align":..., "space_after":..., "size":...}."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        popts = {}
        if isinstance(para, dict):
            popts = para
            runs = para["runs"]
        elif isinstance(para, str):
            runs = [para]
        else:
            runs = para
        p.alignment = popts.get("align", align)
        p.space_after = Pt(popts.get("space_after", space_after))
        if line_spacing:
            p.line_spacing = line_spacing
        for run in runs:
            if isinstance(run, str):
                t, o = run, {}
            else:
                t, o = run
            r = p.add_run()
            r.text = t
            _style(r, size=o.get("size", popts.get("size", size)), color=o.get("color", color),
                   bold=o.get("bold", bold), italic=o.get("italic", False), font=o.get("font", font))
    return tb


def pill(slide, x, y, w, label, color=CYAN, fill=PILL, h=0.34, size=11):
    rect(slide, x, y, w, h, fill=fill, radius=0.08)
    text(slide, x, y, w, h, [[(label, {"bold": True, "color": color})]], size=size,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def checks(slide, x, y, w, h, items, size=13, mark="✓", mark_color=CYAN, color=WHITE, space=7):
    paras = [[(f"{mark}  ", {"color": mark_color, "bold": True}), (it, {"color": color})] for it in items]
    return text(slide, x, y, w, h, paras, size=size, space_after=space)


def code(slide, x, y, w, h, src, size=12, label="PYTHON"):
    rect(slide, x, y, w, h, fill=CODE_BG, radius=0.08, line=LINE)
    if label:
        text(slide, x + w - 1.25, y + 0.08, 1.1, 0.25, [[(label, {"bold": True, "color": DIM})]],
             size=10, align=PP_ALIGN.RIGHT)
    lines = src.strip("\n").splitlines()
    longest = max(len(ln) for ln in lines)
    # 0.6 em per character covers Consolas, Menlo and DejaVu Sans Mono
    fit = (w - 0.45) * 72 / (0.6 * longest)
    if fit < size:
        size = max(10, int(fit * 2) / 2)
        if fit < 10:
            print(f"  code too wide: {longest} chars in {w}in box -> '{max(lines, key=len).strip()[:50]}'")
    paras = []
    for ln in lines:
        stripped = ln.lstrip()
        col = CODE_CMT if stripped.startswith("#") else CODE_FG
        if stripped.startswith("$"):
            col = CYAN
        if stripped.startswith("→"):
            col = GREEN
        paras.append({"runs": [(ln if ln else " ", {"color": col})], "space_after": 0})
    text(slide, x + 0.2, y + 0.18, w - 0.4, h - 0.3, paras, size=size, font=MONO, line_spacing=1.05)


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def table(slide, x, y, w, rows, col_w, row_h=0.42, size=12, header_fill=PILL, header_color=CYAN,
          fills=(CARD, CARD2), first_col_color=WHITE, color=MUTED, col_colors=None):
    nrows, ncols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = gf.table
    gf._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId")).text = \
        "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if i == 0 else fills[(i - 1) % 2])
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            if isinstance(val, tuple):
                r.text, cc = val
            else:
                r.text, cc = val, None
            if i == 0:
                _style(r, size=size, color=header_color, bold=True)
            else:
                c = cc or (col_colors[j] if col_colors and col_colors[j] else (first_col_color if j == 0 else color))
                _style(r, size=size, color=c, bold=(j == 0))
    return tbl


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt.strip()


# ---------------------------------------------------------------- charts (dark styling)
def _dark_chart(chart, size=11):
    chart.has_title = False
    chart.font.size = Pt(size)
    chart.font.name = BODY
    chart.font.color.rgb = rgb(MUTED)


def _axis_dark(ax, grid=True):
    ax.format.line.color.rgb = rgb(LINE)
    ax.tick_labels.font.color.rgb = rgb(MUTED)
    ax.tick_labels.font.size = Pt(10)
    if grid:
        ax.has_major_gridlines = True
        ax.major_gridlines.format.line.color.rgb = rgb(LINE)
        ax.major_gridlines.format.line.width = Pt(0.5)
    else:
        ax.has_major_gridlines = False


# ---------------------------------------------------------------- slide frames
def base(title=None, sub=None, tag=None):
    s = prs.slides.add_slide(BLANK)
    _slide_no[0] += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if title:
        bar_h = 0.9 if sub else 0.58
        rect(s, 0.5, 0.42, 0.07, bar_h, fill=ACCENT, radius=0)
        text(s, 0.7, 0.33, 11.5, 0.65, [[(title, {"bold": True, "color": WHITE})]], size=30)
        if sub:
            text(s, 0.7, 0.93, 11.8, 0.4, sub, size=14, color=MUTED)
    text(s, 0.5, 7.05, 6, 0.3, "ADVANCED ML IN PRODUCTION  ·  DAY 1", size=10, color=DIM)
    right = f"{tag}   |   {_slide_no[0]}" if tag else f"{_slide_no[0]}"
    text(s, 6.8, 7.05, 6.03, 0.3, right, size=10, color=DIM, align=PP_ALIGN.RIGHT)
    return s


def divider(num, title, minutes, window, outcomes, subs, color, tag):
    s = base(tag=tag)
    pill(s, 0.8, 1.7, 1.7, f"MODULE {num}", color=color)
    text(s, 0.8, 2.2, 6.4, 1.8, [[(title, {"bold": True, "color": WHITE})]], size=40, line_spacing=0.95)
    text(s, 0.8, 4.15, 6.4, 0.4, [[(f"{minutes} minutes", {"bold": True, "color": color}),
                                   (f"   |   {window}   |   Outcomes {outcomes}", {"color": MUTED})]], size=16)
    rect(s, 7.6, 1.4, 5.2, 4.7, fill=CARD)
    text(s, 7.95, 1.65, 4.6, 0.35, [[("IN THIS MODULE", {"bold": True, "color": DIM})]], size=11)
    yy = 2.15
    for i, (label, mins) in enumerate(subs, 1):
        rect(s, 7.95, yy, 4.5, 0.82, fill=CARD2, radius=0.08)
        oval(s, 8.1, yy + 0.18, 0.46, color, str(i), color=BG, size=14)
        text(s, 8.75, yy + 0.08, 2.95, 0.66, [[(label, {"color": WHITE, "bold": True})]], size=13,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, 11.6, yy + 0.08, 0.75, 0.66, [[(f"{mins}m", {"color": color, "bold": True})]], size=13,
             anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
        yy += 0.95
    return s


def quiz_checkpoint(tag, checkpoint, check_items, qa, color):
    s = base("Checkpoint & Quick Quiz", "Confirm everyone can move on before starting the next module", tag)
    rect(s, 0.5, 1.55, 4.6, 5.25, fill=CARD)
    pill(s, 0.8, 1.8, 1.6, "CHECKPOINT", color=GREEN)
    text(s, 0.8, 2.3, 4.0, 1.3, [[(checkpoint, {"bold": True, "color": WHITE})]], size=16)
    checks(s, 0.8, 3.65, 4.1, 3.0, check_items, size=13, mark="☐", mark_color=GREEN)
    yy = 1.55
    hq = (5.25 - 0.2 * (len(qa) - 1)) / len(qa)
    for i, (q, a) in enumerate(qa, 1):
        rect(s, 5.4, yy, 7.43, hq, fill=CARD)
        oval(s, 5.6, yy + 0.2, 0.5, color, f"Q{i}", color=BG, size=12)
        text(s, 6.3, yy + 0.14, 6.35, hq - 0.2, [
            [(q, {"bold": True, "color": WHITE})],
            [("Answer: ", {"bold": True, "color": GREEN}), (a, {"color": MUTED})],
        ], size=13, space_after=5)
        yy += hq + 0.2
    return s


# ================================================================ SLIDES
T_INTRO, T_M1, T_M2, T_M3, T_M4 = ("INTRO · 0:00–0:15", "MODULE 1 · 0:15–1:10", "MODULE 2 · 1:10–2:05",
                                   "MODULE 3 · 2:05–3:00", "MODULE 4 · 3:00–4:00")

# ---------- 1. Title
s = base()
pill(s, 5.07, 1.75, 3.2, "DAY 1  ·  FULL PROGRAMME")
text(s, 1.0, 2.25, 11.33, 1.9, [[("Advanced Machine Learning", {"bold": True, "color": WHITE})],
                                [("in Production", {"bold": True, "color": WHITE})]],
     size=48, align=PP_ALIGN.CENTER, space_after=0, line_spacing=0.95)
text(s, 1.0, 4.25, 11.33, 0.5, "From leakage-free pipelines to deployed models: one loan-default case study, end to end.",
     size=17, align=PP_ALIGN.CENTER)
rect(s, 3.67, 4.95, 6.0, 0.55, fill=CARD, radius=0.08)
text(s, 3.67, 4.95, 6.0, 0.55, [[("Day 1   |   4 Hours   |   240 Minutes   |   4 Modules", {"bold": True, "color": CYAN})]],
     size=14, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.0, 5.85, 11.33, 0.4, "Trainer: Muhamad Majed Tuah Bin Mohd Hafizuddin   ·   23 October 2026",
     size=12, color=DIM, align=PP_ALIGN.CENTER)
notes(s, """
0:00 - Welcome (4 min for introductions starts here).
Introduce yourself, then ask each participant: name, role, and one model they have built that never made it to production.
Set the tone: today is hands-on. One case study - a loan-default model - travels through the whole production lifecycle in four hours.
""")

# ---------- 2. Goals
s = base("Programme Goals & Objectives", tag=T_INTRO)
rect(s, 0.5, 1.45, 6.0, 5.35, fill=CARD)
pill(s, 0.8, 1.75, 1.2, "MISSION")
text(s, 0.8, 2.25, 5.4, 0.5, [[("Core Purpose", {"bold": True, "color": WHITE})]], size=22)
text(s, 0.8, 2.85, 5.4, 2.6, [
    "Most organisations already use AI, yet fewer than half of AI projects reach production.",
    "This programme closes the gap between a model that works in a notebook and a service that stays "
    "accurate, explainable and monitored after go-live.",
    "Every exercise uses open-source tools you can adopt at no licence cost.",
], size=15, space_after=10)
rect(s, 0.8, 5.85, 5.4, 0.65, fill=CARD2, radius=0.08)
text(s, 0.95, 5.85, 5.2, 0.65, [[("Key Principle: a model only creates value once it can be trusted in production.",
                                  {"bold": True, "color": CYAN})]], size=13, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 6.83, 1.45, 6.0, 5.35, fill=CARD)
pill(s, 7.13, 1.75, 1.4, "OUTCOMES", color=INDIGO)
text(s, 7.13, 2.25, 5.4, 0.5, [[("Key Learning Outcomes", {"bold": True, "color": WHITE})]], size=22)
text(s, 7.13, 2.8, 5.4, 0.35, "By the end of these 4 hours, you will be able to:", size=13)
checks(s, 7.13, 3.2, 5.55, 3.5, [
    "(a) Choose classic ML, deep learning or an LLM",
    "(b) Detect leakage; build a leakage-free pipeline",
    "(c) Organise training into reproducible scripts",
    "(d) Tune with Optuna and track runs with MLflow",
    "(e) Explain predictions with SHAP in plain language",
    "(f) Detect drift with PSI and KS; set alert thresholds",
    "(g) Serve a model with FastAPI inside Docker",
    "(h) Monitor in production; map the flow to Azure ML",
], size=13, space=4)
notes(s, """
0:04 - Programme overview and objectives (3 min, together with ground rules).
Read the mission. Point out that each outcome letter (a)-(h) maps to a module; checkpoints at the end of each module test them.
Ground rules: cameras/attention on, ask questions any time, help your neighbour during hands-on, phones silent.
""")

# ---------- 3. Roadmap
s = base("Your 4-Hour Roadmap", "Four modules, one case study, a checkpoint at the end of every block", T_INTRO)
blocks = [
    ("0:00–0:15", "Welcome & Ice Breaker", 15, CYAN,
     ["Introductions", "Objectives & ground rules", "\"Would you ship it?\"", "Environment check"],
     "First notebook cell runs"),
    ("0:15–1:10", "Trustworthy Data & Leakage-Free Pipelines", 55, GREEN,
     ["ML vs DL vs LLM", "Data quality & leakage", "Hands-on: find the leaks", "Pipelines & time splits"],
     "Leak-free pipeline on a time split"),
    ("1:10–2:05", "Production Training, Tuning & Tracking", 55, ORANGE,
     ["Notebook to scripts", "Grid / random / Bayesian", "Hands-on: Optuna", "MLflow tracking"],
     "Best trial visible in MLflow"),
    ("2:05–3:00", "Explainability & Drift Monitoring", 55, INDIGO,
     ["SHAP global & local", "Business narrative", "PSI & KS drift tests", "Anomalies & thresholds"],
     "Explanation + PSI table"),
    ("3:00–4:00", "Serving, Monitoring & Cloud", 60, RED,
     ["FastAPI + Docker", "Six-panel dashboard", "Azure ML demo", "Action plan & wrap-up"],
     "curl returns a prediction"),
]
xx = 0.5
total_w = 12.33
for _, _, mins, col, _, _ in blocks:
    seg = total_w * mins / 240
    rect(s, xx, 1.55, seg - 0.04, 0.16, fill=col, radius=0)
    xx += seg
cw, gap = 2.3, 0.2075
for i, (win, title, mins, col, items, ck) in enumerate(blocks):
    x0 = 0.5 + i * (cw + gap)
    rect(s, x0, 1.95, cw, 4.85, fill=CARD)
    text(s, x0 + 0.18, 2.1, cw - 0.36, 0.3, [[(win, {"bold": True, "color": col})]], size=13)
    text(s, x0 + 0.18, 2.45, cw - 0.36, 0.95, [[(title, {"bold": True, "color": WHITE})]], size=15, line_spacing=0.95)
    pill(s, x0 + 0.18, 3.42, 0.95, f"{mins} MIN", color=col, h=0.3, size=10)
    checks(s, x0 + 0.18, 3.9, cw - 0.3, 1.9, items, size=12, mark="›", mark_color=col, space=5)
    rect(s, x0 + 0.12, 5.85, cw - 0.24, 0.8, fill=CARD2, radius=0.08)
    text(s, x0 + 0.22, 5.88, cw - 0.4, 0.75, [[("CHECKPOINT", {"bold": True, "color": GREEN, "size": 9})],
                                              [(ck, {"color": WHITE})]], size=11, space_after=1,
         anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Walk the roadmap left to right in under a minute. The coloured bar on top is proportional to time.
Emphasise the checkpoints: nobody moves to the next module until the checkpoint works - pair up if someone is stuck.
Timing note: the session plan allocates all 240 minutes to content. If the venue requires a comfort break, take 5 minutes from the Module 2 Optuna hands-on and 5 from the Module 4 dashboard walkthrough.
""")

# ---------- 4. Why this matters
s = base("Why This Programme Exists", "Adoption is high, but production is the bottleneck", T_INTRO)
stats = [
    ("78%", "of organisations used AI in 2024, up from 55% a year earlier", "Stanford AI Index, 2025", CYAN),
    ("< 50%", "of AI projects reach production; prototype to production takes ~8 months", "Gartner via Wilkinson, 2024", ORANGE),
    ("294", "published papers across 17 fields were affected by data leakage", "Kapoor & Narayanan, 2023", RED),
    ("63%", "of employers say skill gaps are the biggest barrier to transformation", "World Economic Forum, 2025", INDIGO),
    ("59%", "of Malaysian organisations have deployed or are piloting AI", "Aon, 2026", GREEN),
    ("5%", "of Malaysian organisations are confident they can hire enough AI talent", "Aon, 2026", RED),
]
for i, (big, desc, src, col) in enumerate(stats):
    r_, c_ = divmod(i, 3)
    x0, y0 = 0.5 + c_ * 4.18, 1.55 + r_ * 2.35
    rect(s, x0, y0, 3.98, 2.15, fill=CARD)
    text(s, x0 + 0.25, y0 + 0.15, 3.5, 0.85, [[(big, {"bold": True, "color": col})]], size=44)
    text(s, x0 + 0.25, y0 + 1.0, 3.5, 0.8, desc, size=13, color=WHITE)
    text(s, x0 + 0.25, y0 + 1.75, 3.5, 0.3, src, size=10, color=DIM)
rect(s, 0.5, 6.25, 12.33, 0.6, fill=CARD2, radius=0.08)
text(s, 0.75, 6.25, 11.9, 0.6, [[("Sculley et al. (2015): ", {"bold": True, "color": CYAN}),
                                 ("real-world ML systems carry large hidden technical debt. Today we practise the habits that keep that debt small.",
                                  {"color": WHITE})]], size=13, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Use two numbers only if short on time: fewer than half of projects reach production, and leakage silently broke 294 published papers.
Bridge: 'The models were not bad; the pipelines around them were. That is what we fix today.'
""")

# ---------- 5. Ice breaker
s = base("Ice Breaker: Would You Ship It?", "Poll: vote A, B or C before the reveal", T_INTRO)
rect(s, 0.5, 1.55, 5.6, 5.25, fill=CARD)
pill(s, 0.8, 1.8, 1.4, "SCENARIO", color=ORANGE)
text(s, 0.8, 2.3, 5.0, 1.0, [[("99% accuracy", {"bold": True, "color": WHITE})]], size=48)
text(s, 0.8, 3.3, 5.0, 0.8, "on a loan book where only 1% of borrowers default. Your manager wants it live on Monday.",
     size=15, color=MUTED)
for i, (k, v, col) in enumerate([("A", "Ship it, 99% is excellent", GREEN), ("B", "I need to see more metrics first", ORANGE),
                                 ("C", "Reject it outright", RED)]):
    yy = 4.35 + i * 0.75
    rect(s, 0.8, yy, 5.0, 0.6, fill=CARD2, radius=0.08)
    oval(s, 0.92, yy + 0.09, 0.42, col, k, color=BG, size=13)
    text(s, 1.5, yy, 4.2, 0.6, [[(v, {"color": WHITE})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 6.4, 1.55, 6.43, 5.25, fill=CARD)
pill(s, 6.7, 1.8, 1.3, "REVEAL", color=CYAN)
text(s, 6.7, 2.25, 5.9, 0.7, [[("A model that always says \"no default\" scores 99% and catches zero defaulters.",
                               {"bold": True, "color": WHITE})]], size=15)
labels = [("990", "True negative", CARD2, MUTED), ("0", "False positive", CARD2, MUTED),
          ("10", "Missed defaults", "3A1E2A", RED), ("0", "Caught defaults", "3A1E2A", RED)]
text(s, 6.7, 3.05, 3.0, 0.3, "Per 1,000 loans (always predicts \"no default\")", size=11, color=DIM)
for i, (n, lab, fill, col) in enumerate(labels):
    r_, c_ = divmod(i, 2)
    x0, y0 = 6.7 + c_ * 1.65, 3.4 + r_ * 1.05
    rect(s, x0, y0, 1.55, 0.95, fill=fill, radius=0.08)
    text(s, x0, y0 + 0.05, 1.55, 0.5, [[(n, {"bold": True, "color": col})]], size=22, align=PP_ALIGN.CENTER)
    text(s, x0, y0 + 0.55, 1.55, 0.35, lab, size=10, color=MUTED, align=PP_ALIGN.CENTER)
checks(s, 10.15, 3.4, 2.55, 2.1, ["Accuracy  99%", "Recall  0%", "Business value  none"], size=13,
       mark="•", mark_color=ORANGE, space=10)
rect(s, 6.7, 5.65, 5.85, 0.95, fill=CARD2, radius=0.08)
text(s, 6.85, 5.65, 5.6, 0.95, [[("Seen in our own Session 02 notebook: ", {"bold": True, "color": CYAN}),
                                 ("99.2% accuracy, 0 of 2 test defaults caught. Choose metrics that match the business cost: recall, precision, PR-AUC.",
                                  {"color": WHITE})]], size=12, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
0:07 - Ice breaking (5 min). Run a show-of-hands or chat poll: A, B or C.
Expected answer: B. Ask a 'B' voter what they would want to see (recall, precision, confusion matrix, how the 1% was measured, whether the data leaks).
Reveal: predicting the majority class gets 99% accuracy with zero recall. Mention the Session 02 notebook really produces this result - we will revisit it in Module 1.
Transition: 'Accuracy is the first trap. Leakage is the second, and it is harder to see.'
""")

# ---------- 6. Case study & environment
s = base("One Case Study, the Full Production Lifecycle", "A loan-default model travels through every stage today", T_INTRO)
steps = [
    ("1", "Choose the approach", "ML vs DL vs LLM", GREEN), ("2", "Leakage-free pipeline", "scikit-learn", GREEN),
    ("3", "Modular training", "Python + YAML", ORANGE), ("4", "Tune & track", "Optuna + MLflow", ORANGE),
    ("5", "Explain decisions", "SHAP", INDIGO), ("6", "Watch for drift", "PSI · KS · Isolation Forest", INDIGO),
    ("7", "Serve as an API", "FastAPI + Docker", RED), ("8", "Monitor & scale", "Dashboard + Azure ML", RED),
]
for i, (n, t, tool, col) in enumerate(steps):
    r_, c_ = divmod(i, 4)
    x0, y0 = 0.5 + c_ * 3.13, 1.6 + r_ * 1.55
    rect(s, x0, y0, 2.9, 1.3, fill=CARD)
    oval(s, x0 + 0.2, y0 + 0.38, 0.55, col, n, color=BG, size=16)
    text(s, x0 + 0.9, y0 + 0.2, 1.95, 0.5, [[(t, {"bold": True, "color": WHITE})]], size=14)
    text(s, x0 + 0.9, y0 + 0.7, 1.95, 0.45, [[(tool, {"color": col})]], size=12)
    if c_ < 3:
        arrow(s, x0 + 2.92, y0 + 0.65, x0 + 3.11, y0 + 0.65, color=DIM)
rect(s, 0.5, 4.85, 12.33, 1.95, fill=CARD)
pill(s, 0.8, 5.08, 2.4, "ENVIRONMENT CHECK", color=GREEN)
checks(s, 0.8, 5.55, 5.6, 1.2, ["Python 3.11, Git and Docker Desktop installed",
                                "Repository cloned from the students branch"], size=13, mark_color=GREEN)
code(s, 6.6, 5.05, 6.0, 1.55, """
$ cd advance-ml-fama && jupyter lab
# open day1/01_ai_ml_dl_llm_landscape/
#      ai_ml_landscape_student.ipynb  -> run cell 1
""", size=12, label="TERMINAL")
notes(s, """
0:12 - Environment check (3 min). Everyone opens the student notebook and runs the first cell.
Checkpoint: every participant runs the first notebook cell without error. Helpers: pair anyone who fails with a neighbour; fix during Module 1 lecture time if needed.
Point at the colour coding: green = Module 1, orange = Module 2, indigo = Module 3, red = Module 4.
""")

# ================================================================ MODULE 1
s = divider(1, "Trustworthy Data & Leakage-Free Pipelines", 55, "0:15 – 1:10", "(a) and (b)",
            [("Choosing the right approach", 10), ("Data quality and data leakage", 10),
             ("Hands-on: finding leakage", 15), ("Leakage-free pipelines & validation", 20)], GREEN, T_M1)
notes(s, "0:15 - Module 1 starts. Outcomes (a) choose the approach and (b) detect leakage and build a leakage-free pipeline.")

# ---------- AI hierarchy
s = base("AI, ML, DL and LLMs: Nested, Not Competing", "Each layer is a specialised subset of the one around it", T_M1)
rings = [("ARTIFICIAL INTELLIGENCE", 0.5, 1.55, 5.9, 5.25, CARD, CYAN),
         ("MACHINE LEARNING", 0.85, 2.15, 5.2, 4.35, CARD2, GREEN),
         ("DEEP LEARNING", 1.2, 2.75, 4.5, 3.45, PILL, ORANGE),
         ("LARGE LANGUAGE MODELS", 1.55, 3.35, 3.8, 2.55, "3A3F6E", RED)]
for lab, x0, y0, w0, h0, fill, col in rings:
    rect(s, x0, y0, w0, h0, fill=fill, radius=0.25, line=col)
    text(s, x0 + 0.25, y0 + 0.12, w0 - 0.5, 0.35, [[(lab, {"bold": True, "color": col})]], size=12)
text(s, 1.8, 4.0, 3.3, 1.6, [[("GPT, Llama, Claude", {"color": WHITE, "bold": True})],
                             [("Transformers trained on huge text corpora", {"color": MUTED})]],
     size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
defs = [
    ("AI", "Any system that mimics intelligent behaviour: rules, search, planning or learning.", CYAN),
    ("ML", "Learns patterns from examples. Gradient boosting and random forests dominate tabular data.", GREEN),
    ("DL", "Multi-layer neural networks that learn features from raw images, audio and sequences.", ORANGE),
    ("LLM", "Deep learning models for language: summarise, answer, draft and reason over text.", RED),
]
for i, (k, d, col) in enumerate(defs):
    y0 = 1.55 + i * 1.33
    rect(s, 6.7, y0, 6.13, 1.18, fill=CARD)
    oval(s, 6.9, y0 + 0.29, 0.6, col, k, color=BG, size=13 if len(k) < 3 else 11)
    text(s, 7.7, y0 + 0.1, 4.95, 1.0, d, size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
0:15 - Choosing the right approach (10 min total for this and the next slide).
Key message: these are nested, not rivals. An LLM is a deep learning model, which is a machine learning model, which is AI.
Ask: 'Which layer does our loan-default model belong to?' Answer: ML - tabular data, classic algorithms.
""")

# ---------- Choosing the approach
s = base("Choosing the Right Approach", "Let the shape of the data, not the hype, pick the tool", T_M1)
table(s, 0.5, 1.5, 12.33, [
    ["Business problem", "Data shape", "Best fit", "Why"],
    ["Predict loan default", "Tabular rows", ("Classic ML (gradient boosting)", GREEN), "Accurate, fast, explainable with SHAP"],
    ["Segment customers with no labels", "Tabular, no target", ("Unsupervised ML (clustering)", GREEN), "Finds groups without labelled examples"],
    ["Forecast 7-day electricity demand", "Time series", ("Classic ML / forecasting", GREEN), "Lags and seasonality; cheap to retrain"],
    ["Detect defects on a conveyor camera", "Images", ("Deep learning (CNN)", ORANGE), "Learns visual features automatically"],
    ["Transcribe customer service calls", "Audio", ("Deep learning (speech model)", ORANGE), "Pretrained models handle raw audio"],
    ["Answer staff questions from policy PDFs", "Unstructured text", ("LLM + retrieval (RAG)", RED), "Grounds answers in your own documents"],
], [3.55, 2.1, 3.3, 3.38], row_h=0.58, size=13)
rect(s, 0.5, 5.75, 12.33, 1.05, fill=CARD2, radius=0.08)
text(s, 0.75, 5.75, 11.9, 1.05, [
    [("Rule of thumb: ", {"bold": True, "color": CYAN}),
     ("most enterprise data is tabular, so start with classic ML. Reach for deep learning or LLMs when the data is images, audio or free text.",
      {"color": WHITE})],
    [("Quick poll: ", {"bold": True, "color": ORANGE}),
     ("cover the 'Best fit' column and let the room call each answer.", {"color": MUTED})]], size=13,
     anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Cover the third column and let participants call out the answer for each row (2 minutes).
Discussion prompt: 'Why not use an LLM for loan default?' - tabular data, need for calibrated probabilities, regulatory explainability, cost and latency.
Source activity: day1/01_ai_ml_dl_llm_landscape README, Activity 1 (10 business problems) if you want a longer version.
""")

# ---------- Data quality
s = base("Five Dimensions of Data Quality", "Check them before you trust any metric", T_M1)
dims = [
    ("Integrity", "Schema and types are as expected", "credit_score arrives as text", CYAN),
    ("Completeness", "Missing values are limited and understood", "8.8% of annual_income missing", GREEN),
    ("Accuracy", "Values are correct and realistic", "11.9% of incomes outside IQR fences", ORANGE),
    ("Timeliness", "Data is fresh enough for the decision", "Bureau score pulled 6 months ago", INDIGO),
    ("Validity", "Values respect business rules", "Negative employment_years", RED),
]
for i, (k, d, ex, col) in enumerate(dims):
    x0 = 0.5 + i * 2.5
    rect(s, x0, 1.55, 2.33, 2.75, fill=CARD)
    text(s, x0 + 0.18, 1.7, 2.0, 0.4, [[(k, {"bold": True, "color": col})]], size=16)
    text(s, x0 + 0.18, 2.15, 2.0, 0.9, d, size=13, color=WHITE)
    rect(s, x0 + 0.12, 3.2, 2.09, 0.95, fill=CARD2, radius=0.08)
    text(s, x0 + 0.22, 3.2, 1.9, 0.95, [[("LOAN DATA", {"bold": True, "color": DIM, "size": 9})], [(ex, {"color": MUTED})]],
         size=12, anchor=MSO_ANCHOR.MIDDLE, space_after=1)
text(s, 0.5, 4.5, 6, 0.35, [[("WHY IS IT MISSING?  The answer changes the fix", {"bold": True, "color": DIM})]], size=11)
miss = [
    ("MCAR", "Missing completely at random", "A form field dropped by a system glitch", "Simple imputation is fine"),
    ("MAR", "Depends on other observed columns", "Income blank more often for self-employed", "Impute using the related columns"),
    ("MNAR", "Depends on the missing value itself", "High-debt applicants skip the income field", "Add a was_missing flag; investigate"),
]
for i, (k, d, ex, fix) in enumerate(miss):
    x0 = 0.5 + i * 4.18
    rect(s, x0, 4.9, 3.98, 1.9, fill=CARD)
    text(s, x0 + 0.2, 5.0, 3.6, 1.8, [
        [(k + "  ", {"bold": True, "color": CYAN, "size": 16}), (d, {"color": WHITE})],
        [("e.g. ", {"color": DIM}), (ex, {"color": MUTED})],
        [("Fix: ", {"bold": True, "color": GREEN}), (fix, {"color": MUTED})],
    ], size=13, space_after=6)
notes(s, """
0:25 - Data quality and data leakage (10 min for this and the next slide).
The completeness and accuracy figures are real outputs of the Session 02 notebook's loan dataset.
Ask: 'Which missingness type is most dangerous?' MNAR - the fact that a value is missing carries signal, and naive imputation hides it.
""")

# ---------- Leakage types
s = base("Four Ways Information Leaks Into Training", "Leakage gives optimistic validation scores and poor production performance", T_M1)
leaks = [
    ("Target leakage", RED, "A feature is created after, or because of, the outcome.",
     "days_to_default, post_approval_status_flag", "Keep only features known at decision time"),
    ("Train-test contamination", ORANGE, "Preprocessing is fitted on all rows before the split.",
     "Scaler or imputer fitted on the full dataset", "Split first; fit inside a Pipeline"),
    ("Temporal leakage", CYAN, "A random split lets the model learn from the future.",
     "2025 loans train a model tested on 2023", "Time-ordered split or TimeSeriesSplit"),
    ("Group leakage", INDIGO, "The same entity appears in both train and test.",
     "A repeat customer in both sets", "GroupKFold or split by customer_id"),
]
for i, (k, col, d, ex, fix) in enumerate(leaks):
    x0 = 0.5 + i * 3.13
    rect(s, x0, 1.55, 2.95, 4.35, fill=CARD)
    rect(s, x0 + 0.2, 1.75, 2.55, 0.06, fill=col, radius=0)
    text(s, x0 + 0.2, 1.9, 2.6, 0.45, [[(k, {"bold": True, "color": col})]], size=16)
    text(s, x0 + 0.2, 2.4, 2.6, 1.0, d, size=13, color=WHITE)
    text(s, x0 + 0.2, 3.45, 2.6, 0.3, [[("LOAN EXAMPLE", {"bold": True, "color": DIM})]], size=10)
    text(s, x0 + 0.2, 3.72, 2.6, 0.8, ex, size=12, color=MUTED)
    rect(s, x0 + 0.12, 4.75, 2.71, 0.95, fill=CARD2, radius=0.08)
    text(s, x0 + 0.22, 4.75, 2.5, 0.95, [[("FIX  ", {"bold": True, "color": GREEN, "size": 10})], [(fix, {"color": WHITE})]],
         size=12, anchor=MSO_ANCHOR.MIDDLE, space_after=1)
rect(s, 0.5, 6.1, 12.33, 0.7, fill=CARD2, radius=0.08)
text(s, 0.75, 6.1, 11.9, 0.7, [[("scikit-learn guidance: ", {"bold": True, "color": CYAN}),
                                ("split the data before any preprocessing, and keep every fitted step inside a Pipeline.",
                                 {"color": WHITE})]], size=13, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Spend about 90 seconds per type. For each, ask the room for one example from their own work.
Note: the coloured lines here mirror the Session 05 transparency-spectrum styling.
Transition: 'Now you will hunt for every one of these in a real dataset.'
""")

# ---------- Hands-on: find the leakage
s = base("Hands-On: Find Every Leakage Source", "15 minutes  ·  day1/02_data_quality_leakage/notebook.ipynb", T_M1)
cols = ["customer_id", "loan_amount", "annual_income", "employment_years", "credit_score", "days_since_last_payment",
        "existing_loans", "loan_type", "region", "loan_approval_date", "days_to_default", "post_approval_status_flag"]
rect(s, 0.5, 1.55, 7.2, 5.25, fill=CARD)
text(s, 0.8, 1.75, 6.6, 0.35, [[("THE 12 CANDIDATE COLUMNS", {"bold": True, "color": DIM})]], size=11)
for i, c in enumerate(cols):
    r_, c_ = divmod(i, 2)
    x0, y0 = 0.8 + c_ * 3.35, 2.15 + r_ * 0.62
    rect(s, x0, y0, 3.2, 0.5, fill=CARD2, radius=0.08)
    text(s, x0 + 0.15, y0, 3.0, 0.5, [[(c, {"color": CODE_FG})]], size=13, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.8, 5.95, 6.6, 0.7, [[("Target: ", {"bold": True, "color": ORANGE}),
                              ("defaulted_in_next_30_days  ·  1,200 loans  ·  0.7% default rate", {"color": MUTED})]], size=13)
rect(s, 8.0, 1.55, 4.83, 5.25, fill=CARD)
pill(s, 8.3, 1.8, 1.6, "YOUR TASK", color=ORANGE)
text(s, 8.3, 2.3, 4.3, 1.0, [[("For every column ask: would I know this value at the moment the loan is approved?",
                              {"bold": True, "color": WHITE})]], size=15)
for i, (k, d, col) in enumerate([("SAFE", "Known at decision time", GREEN),
                                 ("RISKY", "Usable only with care (split or as-of date)", ORANGE),
                                 ("LEAK", "Only exists after the outcome", RED)]):
    yy = 3.55 + i * 0.72
    pill(s, 8.3, yy, 1.05, k, color=col, h=0.42, size=12)
    text(s, 9.5, yy, 3.2, 0.42, d, size=13, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
text(s, 8.3, 5.75, 4.3, 0.9, "Then run the leaky and clean models in the notebook and compare their reports.",
     size=13, color=MUTED)
notes(s, """
0:35 - Finding leakage (15 min). Individual work for 8 minutes, then compare with a neighbour for 3 minutes, then reveal (next slide) for 4 minutes.
Known issue for the trainer: in the current notebook one 'leaky model' cell raises a NaN error (the second copy of the cell imputes first and works), and the clean AUC values are below 0.5 because the synthetic data has only about 8 defaults. Use it as a discussion point about rare classes; do not present those AUCs as benchmarks.
""")

# ---------- Answer key
s = base("Answer Key: Every Leakage Source", "Reveal after pairs have compared their answers", T_M1)
key = [
    ("LEAK", RED, [("days_to_default", "Only exists once a default happens"),
                   ("post_approval_status_flag", "Recorded after approval, encodes the outcome")]),
    ("RISKY", ORANGE, [("loan_approval_date", "Not a feature; use it to order the split by time"),
                       ("customer_id", "Not a feature; use it to group repeat customers"),
                       ("days_since_last_payment", "Safe only if measured as-of the application date"),
                       ("annual_income · credit_score", "Have gaps; impute inside the Pipeline, not before")]),
    ("SAFE", GREEN, [("loan_amount · employment_years", "Known when the application is submitted"),
                     ("existing_loans · loan_type · region", "Known when the application is submitted")]),
]
yy = 1.5
for lab, col, items in key:
    hh = 0.2 + 0.54 * len(items)
    rect(s, 0.5, yy, 12.33, hh, fill=CARD)
    pill(s, 0.75, yy + hh / 2 - 0.22, 1.2, lab, color=col, h=0.44, size=13)
    for j, (c, why) in enumerate(items):
        y1 = yy + 0.1 + j * 0.54
        rect(s, 2.25, y1, 4.4, 0.46, fill=CARD2, radius=0.08)
        text(s, 2.4, y1, 4.2, 0.46, [[(c, {"color": col})]], size=13, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 6.9, y1, 5.8, 0.46, why, size=13, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    yy += hh + 0.14
notes(s, """
Reveal one band at a time. The RISKY band generates the best discussion: dates and IDs are not leaks in themselves, but using them as features or ignoring them in the split creates temporal and group leakage.
Ask: 'How would you prove days_since_last_payment is safe?' - check the data dictionary for its as-of timestamp.
""")

# ---------- Wrong vs right flow
s = base("The Leakage-Free Pattern", "Same data, same model: only the order of operations changes", T_M1)


def flow(y0, label, col, steps_, verdict):
    rect(s, 0.5, y0, 12.33, 2.35, fill=CARD)
    pill(s, 0.8, y0 + 0.22, 1.3, label, color=col)
    w0, g0 = 2.15, 0.42
    for i, st in enumerate(steps_):
        x0 = 0.8 + i * (w0 + g0)
        rect(s, x0, y0 + 0.75, w0, 0.85, fill=CARD2, radius=0.08)
        text(s, x0 + 0.08, y0 + 0.75, w0 - 0.16, 0.85, [[(st, {"color": WHITE})]], size=13,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(steps_) - 1:
            arrow(s, x0 + w0 + 0.04, y0 + 1.17, x0 + w0 + g0 - 0.04, y0 + 1.17, color=col)
    text(s, 0.8, y0 + 1.75, 11.7, 0.45, [[(verdict, {"color": col, "bold": True})]], size=13)


flow(1.55, "WRONG", RED, ["Full dataset", "Fit imputer & scaler on all rows", "Random split", "Train model", "Evaluate"],
     "Test-set statistics and future loans shape training, so the score is optimistic and production disappoints.")
flow(4.15, "RIGHT", GREEN, ["Full dataset", "Time-ordered split (oldest 80%)", "Fit Pipeline on train only",
                            "Tune on validation folds", "Evaluate once on newest 20%"],
     "The holdout looks like next month's applicants, so the score estimates real production performance.")
text(s, 0.5, 6.6, 12.33, 0.35, "Kapoor & Narayanan (2023); scikit-learn developers, Common pitfalls and recommended practices.",
     size=10, color=DIM)
notes(s, """
0:50 - Leakage-free pipelines and validation (20 min for this slide and the next three).
Point out that both rows use exactly the same steps; only the order and the split strategy differ.
""")

# ---------- Pipeline code
s = base("Building the Pipeline in scikit-learn", "Every learned statistic lives inside one object fitted on training rows only", T_M1)
code(s, 0.5, 1.5, 7.6, 5.3, """
df = df.sort_values("loan_approval_date")
cut = int(len(df) * 0.8)              # oldest 80% = train
train, test = df.iloc[:cut], df.iloc[cut:]

num = ["loan_amount", "annual_income", "employment_years",
       "credit_score", "days_since_last_payment", "existing_loans"]
cat = ["loan_type", "region"]

num_pipe = Pipeline([("impute", SimpleImputer(strategy="median")),
                     ("scale", StandardScaler())])
cat_pipe = Pipeline([("impute", SimpleImputer(strategy="most_frequent")),
                     ("onehot", OneHotEncoder(handle_unknown="ignore"))])
pre = ColumnTransformer([("num", num_pipe, num), ("cat", cat_pipe, cat)])

model = Pipeline([("pre", pre),
                  ("clf", LogisticRegression(max_iter=1000))])
model.fit(train[num + cat], train["defaulted_in_next_30_days"])
""", size=12)
why = [("fit() learns from train only", "Medians, scales and categories never see the test rows.", GREEN),
       ("One object, every stage", "The same transforms run in validation, batch scoring and the API.", CYAN),
       ("One artefact to ship", "joblib.dump(model) saves preprocessing and model together.", ORANGE)]
for i, (k, d, col) in enumerate(why):
    y0 = 1.5 + i * 1.82
    rect(s, 8.4, y0, 4.43, 1.65, fill=CARD)
    text(s, 8.65, y0 + 0.18, 4.0, 1.35, [[(k, {"bold": True, "color": col})], [(d, {"color": WHITE})]], size=14, space_after=6)
notes(s, """
Live-code or walk through this cell from the Session 02 notebook ('Full loan default workflow with temporal split').
Emphasise handle_unknown='ignore': new loan types in production will not crash the pipeline.
Ask: 'Where would target encoding go?' Inside the Pipeline too, otherwise it is train-test contamination.
""")

# ---------- Validation strategies
s = base("Pick a Validation Strategy That Fits the Data", "Train on blue, score on orange", T_M1)
strat = [
    ("Stratified K-Fold", "Independent rows; keeps the rare default rate in every fold.", "kfold", CYAN),
    ("TimeSeriesSplit", "Ordered data: always train on the past, test on the future.", "time", GREEN),
    ("GroupKFold", "Repeat entities: all rows of one customer stay in one fold.", "group", INDIGO),
]
for i, (k, d, kind, col) in enumerate(strat):
    x0 = 0.5 + i * 4.18
    rect(s, x0, 1.55, 3.98, 5.25, fill=CARD)
    text(s, x0 + 0.25, 1.75, 3.5, 0.45, [[(k, {"bold": True, "color": col})]], size=18)
    text(s, x0 + 0.25, 2.25, 3.5, 0.8, d, size=13, color=WHITE)
    groups = list("AAABBCCCDD")
    for f in range(4):
        y1 = 3.3 + f * 0.62
        text(s, x0 + 0.2, y1, 0.6, 0.4, f"Fold {f + 1}", size=10, color=DIM, anchor=MSO_ANCHOR.MIDDLE)
        for b in range(10):
            if kind == "time":
                train_n = 2 * (f + 1)
                fill = ACCENT if b < train_n else (ORANGE if b < train_n + 2 else LINE)
            elif kind == "kfold":
                fill = ORANGE if b in (f * 2 + 1, (f * 2 + 6) % 10) else ACCENT
            else:
                test_group = "ABCD"[f]
                fill = ORANGE if groups[b] == test_group else ACCENT
            rect(s, x0 + 0.85 + b * 0.29, y1 + 0.05, 0.24, 0.3, fill=fill, radius=0.04)
            if kind == "group" and f == 0:
                text(s, x0 + 0.85 + b * 0.29, 3.0, 0.24, 0.28, groups[b], size=10, color=MUTED, align=PP_ALIGN.CENTER)
    use = {"kfold": "Use for: cross-sectional snapshots, rare classes",
           "time": "Use for: credit, fraud, demand, anything that predicts the future",
           "group": "Use for: customers with many loans, patients, devices"}[kind]
    rect(s, x0 + 0.15, 5.9, 3.68, 0.75, fill=CARD2, radius=0.08)
    text(s, x0 + 0.27, 5.9, 3.45, 0.75, use, size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
For the loan case: TimeSeriesSplit for tuning (applications arrive over time) plus a final holdout of the newest 20%. If customers repeat, combine time ordering with grouping (no customer in both sides).
Demo: the 'Visualise TimeSeriesSplit folds' cell in the Session 02 notebook draws this exact picture.
""")

# ---------- M1 checkpoint
s = quiz_checkpoint(T_M1, "Your pipeline is scored on a time-ordered split with no leaked feature.",
                    ["Leaky columns removed", "Split happens before any preprocessing",
                     "All preprocessing inside the Pipeline", "Metric reported on the newest 20% of loans"],
                    [("Why is fitting a scaler on the full dataset a problem?",
                      "Test-set statistics leak into training, so the validation score is optimistic."),
                     ("When should you prefer a time-based split?",
                      "Whenever the model predicts future events from past data: credit, fraud, demand."),
                     ("Which loan column is pure target leakage?",
                      "days_to_default: it only exists after the borrower has defaulted.")], GREEN)
notes(s, """
1:05 - Module 1 checkpoint (last 5 minutes of the module). Walk the room; each participant shows the four ticks.
Ask the quiz questions aloud; cold-call gently.
""")

# ================================================================ MODULE 2
s = divider(2, "Production Training, Tuning & Tracking", 55, "1:10 – 2:05", "(c) and (d)",
            [("From notebook to modular scripts", 15), ("Hyperparameter optimisation with Optuna", 20),
             ("Experiment tracking with MLflow", 20)], ORANGE, T_M2)
notes(s, "1:10 - Module 2 starts. Outcomes (c) modular, reproducible scripts and (d) Optuna tuning with MLflow tracking.")

# ---------- Modular architecture
s = base("From Notebook to Modular Scripts", "day1/03_production_training/scripts: one responsibility per file", T_M2)
rect(s, 0.5, 1.6, 2.4, 1.2, fill=CARD, line=ORANGE)
text(s, 0.5, 1.6, 2.4, 1.2, [[("config.yaml", {"bold": True, "color": ORANGE, "font": MONO})],
                             [("data · model · paths", {"color": MUTED})]], size=13, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 2.95, 2.2, 4.0, 2.2, color=ORANGE)
rect(s, 4.05, 1.6, 5.2, 1.2, fill=CARD2, line=CYAN)
text(s, 4.05, 1.6, 5.2, 1.2, [[("main.py", {"bold": True, "color": CYAN, "font": MONO})],
                              [("Orchestrator: reads config, sets logging, runs each step", {"color": WHITE})]], size=13,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
mods = [("data_loader.py", "Load CSV, validate target, split"), ("features.py", "Impute and one-hot encode"),
        ("train.py", "Build model from config; save"), ("evaluate.py", "Metrics to JSON report")]
for i, (m, d) in enumerate(mods):
    x0 = 0.5 + i * 3.13
    rect(s, x0, 3.35, 2.95, 1.3, fill=CARD)
    text(s, x0 + 0.15, 3.35, 2.65, 1.3, [[(m, {"bold": True, "color": WHITE, "font": MONO})], [(d, {"color": MUTED})]],
         size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, 6.65, 2.82, x0 + 1.47, 3.33, color=DIM, width=1.25)
arts = [("model.joblib", "Fitted model"), ("metrics.json", "Accuracy, precision, recall, F1"),
        ("training.log", "Timestamped INFO log")]
text(s, 0.5, 4.9, 6, 0.35, [[("ARTIFACTS WRITTEN TO artifacts/", {"bold": True, "color": DIM})]], size=11)
for i, (a, d) in enumerate(arts):
    x0 = 0.5 + i * 2.6
    rect(s, x0, 5.3, 2.45, 1.0, fill=CARD2, radius=0.08)
    text(s, x0 + 0.1, 5.3, 2.25, 1.0, [[(a, {"bold": True, "color": GREEN, "font": MONO})], [(d, {"color": MUTED})]],
         size=12, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 8.45, 4.9, 4.38, 1.9, fill=CARD)
checks(s, 8.7, 5.05, 4.0, 1.7, ["Each file is testable on its own", "Changes are reviewable in Git",
                                "A scheduler can run it without Jupyter"], size=13, mark_color=ORANGE)
notes(s, """
1:10 - From notebook to modular scripts (15 min for this slide and the next).
Open each file briefly in the IDE. Ask: 'If the data source moves from CSV to a database, which file changes?' Only data_loader.py.
""")

# ---------- Reproducibility
s = base("Configuration, Seeds and Logging", "Run the same command twice and get the same model", T_M2)
code(s, 0.5, 1.5, 5.4, 3.75, """
data:
  path: data/production_training_sample.csv
  target_column: churn
  test_size: 0.2
  random_state: 42
model:
  type: logistic_regression
  random_state: 42
  max_iter: 500
paths:
  model_path: artifacts/model.joblib
  report_path: artifacts/metrics.json
  log_path: artifacts/training.log
""", size=12, label="config.yaml")
code(s, 6.2, 1.5, 6.63, 2.05, """
# run from the repository root
$ python day1/03_production_training/scripts/main.py \\
    --config day1/03_production_training/scripts/config.yaml
→ ... | INFO | Starting training workflow
→ ... | INFO | Training complete. Metrics: {...}
""", size=12, label="TERMINAL")
princ = [("Config, not code", "Change behaviour without editing Python", ORANGE),
         ("Fixed seeds", "random_state: 42 everywhere", CYAN),
         ("Structured logs", "Timestamp | level | message", GREEN),
         ("Saved artefacts", "Model + metrics + log per run", INDIGO)]
for i, (k, d, col) in enumerate(princ):
    r_, c_ = divmod(i, 2)
    x0, y0 = 6.2 + c_ * 3.36, 3.75 + r_ * 0.8
    rect(s, x0, y0, 3.27, 0.7, fill=CARD, radius=0.08)
    text(s, x0 + 0.15, y0, 3.0, 0.7, [[(k + "  ", {"bold": True, "color": col}), (d, {"color": MUTED, "size": 11})]],
         size=13, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 0.5, 5.5, 12.33, 1.3, fill=CARD2, radius=0.08)
pill(s, 0.75, 5.7, 1.5, "HANDS-ON", color=ORANGE)
text(s, 2.45, 5.55, 10.2, 1.2, [[("Run main.py once, then change ", {"color": WHITE}),
                                 ("test_size", {"color": CYAN, "font": MONO}), (" or ", {"color": WHITE}),
                                 ("max_iter", {"color": CYAN, "font": MONO}),
                                 (" in config.yaml and run again. Compare the two metrics.json files. Same config, same seed: identical result.",
                                  {"color": WHITE})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Important: always pass --config explicitly. The default path in main.py is Windows-style and will not resolve on macOS/Linux.
Paths inside config.yaml are relative to the repository root, so run the command from the root.
If data/production_training_sample.csv is missing, data_loader.py generates a synthetic dataset automatically.
""")

# ---------- Params vs hyperparams
s = base("Parameters vs Hyperparameters", "What the model learns versus what you decide", T_M2)
for i, (k, sub, items, col) in enumerate([
    ("Model parameters", "Learned from data during fit()", ["Logistic regression coefficients",
                                                           "Tree split thresholds", "Neural network weights"], CYAN),
    ("Hyperparameters", "Set by you before training", ["n_estimators, max_depth", "learning_rate, C (regularisation)",
                                                      "min_samples_split"], ORANGE)]):
    x0 = 0.5 + i * 6.33
    rect(s, x0, 1.55, 6.0, 3.3, fill=CARD)
    text(s, x0 + 0.3, 1.75, 5.4, 0.5, [[(k, {"bold": True, "color": col})]], size=22)
    text(s, x0 + 0.3, 2.3, 5.4, 0.4, sub, size=14, color=MUTED)
    checks(s, x0 + 0.3, 2.9, 5.4, 1.9, items, size=15, mark_color=col, space=9)
rect(s, 0.5, 5.1, 7.6, 1.7, fill=CARD)
text(s, 0.8, 5.25, 7.1, 1.45, [[("Baking analogy", {"bold": True, "color": WHITE})],
                               [("The oven temperature and time are hyperparameters; you set them. How the ingredients "
                                 "combine inside is learned. Set them badly and the cake burns: the model overfits.", {"color": MUTED})]],
     size=14, space_after=6)
rect(s, 8.4, 5.1, 4.43, 1.7, fill=CARD2)
text(s, 8.65, 5.2, 4.0, 1.5, [[("Golden rule", {"bold": True, "color": RED})],
                              [("Tune on validation folds. Never let the test set choose hyperparameters.", {"color": WHITE})]],
     size=14, anchor=MSO_ANCHOR.MIDDLE, space_after=6)
notes(s, """
1:25 - Hyperparameter optimisation (20 min for this and the next two slides, including 10 min hands-on).
Quick check: ask the room to classify 'number of trees' (hyper) and 'feature weight' (parameter).
""")

# ---------- Search strategies (scatter charts)
s = base("Grid, Random and Bayesian Search", "25 trials each over max_depth (3–20) and n_estimators (50–300)", T_M2)
rng = np.random.default_rng(7)
grid_pts = [(d, n) for d in np.linspace(3, 20, 5) for n in np.linspace(50, 300, 5)]
rand_pts = list(zip(rng.uniform(3, 20, 25), rng.uniform(50, 300, 25)))
bayes = list(zip(rng.uniform(3, 20, 8), rng.uniform(50, 300, 8)))
for k in range(17):
    sd = 1 - k / 20
    bayes.append((float(np.clip(rng.normal(12.5, 3.2 * sd), 3, 20)), float(np.clip(rng.normal(225, 45 * sd), 50, 300))))
panels = [("Grid search", grid_pts, MUTED, "Tests every combination. Wastes trials on unimportant dimensions; cost explodes with more hyperparameters."),
          ("Random search", rand_pts, CYAN, "25 distinct values per axis. Bergstra & Bengio (2012): as good or better than grid for a fraction of the budget."),
          ("Bayesian (Optuna TPE)", bayes, ORANGE, "Learns from past trials and concentrates new ones where scores are high.")]
for i, (k, pts, col, d) in enumerate(panels):
    x0 = 0.5 + i * 4.18
    rect(s, x0, 1.55, 3.98, 5.25, fill=CARD)
    text(s, x0 + 0.25, 1.7, 3.5, 0.4, [[(k, {"bold": True, "color": col})]], size=17)
    cd = XyChartData()
    ser = cd.add_series("trials")
    for a, b in pts:
        ser.add_data_point(round(float(a), 2), round(float(b), 1))
    ch = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(x0 + 0.1), Inches(2.1), Inches(3.78), Inches(2.9), cd).chart
    _dark_chart(ch, 10)
    ch.has_legend = False
    ser0 = ch.plots[0].series[0]
    ser0.marker.style = XL_MARKER_STYLE.CIRCLE
    ser0.marker.size = 7
    ser0.marker.format.fill.solid()
    ser0.marker.format.fill.fore_color.rgb = rgb(col)
    ser0.marker.format.line.fill.background()
    ser0.format.line.fill.background()
    xa, ya = ch.category_axis, ch.value_axis
    xa.minimum_scale, xa.maximum_scale = 0, 22
    ya.minimum_scale, ya.maximum_scale = 0, 320
    xa.major_unit, ya.major_unit = 5, 100
    _axis_dark(xa, grid=False)
    _axis_dark(ya)
    text(s, x0 + 0.25, 5.05, 3.5, 1.65, d, size=12, color=WHITE)
notes(s, """
Point out the grid's problem: only 5 distinct values of max_depth are ever tried. If only one hyperparameter matters, random search tests 25 distinct values of it with the same budget.
Bayesian search (Optuna's default TPE sampler) spends the first trials exploring, then exploits the promising region - visible as the orange cluster.
Points are simulated for illustration.
""")

# ---------- Optuna
s = base("Hands-On: Tuning With Optuna", "10 minutes  ·  day1/04_hyperparameter_optim/notebook.ipynb", T_M2)
code(s, 0.5, 1.5, 7.6, 5.3, """
import optuna
from sklearn.ensemble import RandomForestClassifier
from sklearn.model_selection import cross_val_score

def objective(trial):
    params = {
        "n_estimators": trial.suggest_int("n_estimators", 50, 300),
        "max_depth": trial.suggest_int("max_depth", 3, 20),
        "min_samples_split": trial.suggest_int("min_samples_split", 2, 15),
    }
    clf = RandomForestClassifier(**params, random_state=42)
    # score on validation folds, never on the test set
    return cross_val_score(clf, X_train, y_train,
                           cv=5, scoring="roc_auc").mean()

study = optuna.create_study(direction="maximize")
study.optimize(objective, n_trials=30)
print(study.best_params, round(study.best_value, 3))
""", size=12)
vocab = [("Trial", "One evaluation of one set of hyperparameters", CYAN),
         ("Study", "The whole optimisation: all trials plus the best", ORANGE),
         ("Sampler", "Chooses the next trial; TPE (Bayesian) by default", GREEN),
         ("Pruner", "Stops hopeless trials early to save compute", INDIGO)]
for i, (k, d, col) in enumerate(vocab):
    y0 = 1.5 + i * 1.0
    rect(s, 8.4, y0, 4.43, 0.88, fill=CARD, radius=0.08)
    text(s, 8.6, y0, 4.1, 0.88, [[(k, {"bold": True, "color": col})], [(d, {"color": MUTED, "size": 12})]], size=14,
         anchor=MSO_ANCHOR.MIDDLE, space_after=1)
rect(s, 8.4, 5.55, 4.43, 1.25, fill=CARD2, radius=0.08)
text(s, 8.6, 5.55, 4.1, 1.25, [[("Stretch: ", {"bold": True, "color": ORANGE}),
                               ("plot_param_importances(study) shows which hyperparameter mattered most.", {"color": WHITE})]],
     size=13, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Participants run the Optuna cells. 30 trials with 5-fold CV on a small dataset takes about a minute on a laptop; reduce n_trials to 15 on slow machines.
While it runs, ask: 'Why cross_val_score and not the test set?' (golden rule from two slides ago).
""")

# ---------- MLflow
s = base("Experiment Tracking With MLflow", "If it is not logged, it did not happen", T_M2)
logs = [("PARAMS", "max_depth=12, n_estimators=220", CYAN), ("METRICS", "test_auc, recall, f1", GREEN),
        ("ARTIFACTS", "model, plots, metrics.json", ORANGE), ("TAGS", "git commit, data version, owner", INDIGO)]
for i, (k, d, col) in enumerate(logs):
    x0 = 0.5 + i * 3.13
    rect(s, x0, 1.55, 2.95, 1.25, fill=CARD)
    text(s, x0 + 0.2, 1.65, 2.6, 1.1, [[(k, {"bold": True, "color": col})], [(d, {"color": WHITE})]], size=13,
         anchor=MSO_ANCHOR.MIDDLE, space_after=4)
code(s, 0.5, 3.05, 7.6, 3.75, """
import mlflow, mlflow.sklearn

mlflow.set_experiment("Bootcamp_Optuna_Experiments")
with mlflow.start_run(run_name="Best_RF_Model"):
    mlflow.log_params(study.best_params)
    best = RandomForestClassifier(**study.best_params, random_state=42)
    best.fit(X_train, y_train)
    auc = roc_auc_score(y_test, best.predict_proba(X_test)[:, 1])
    mlflow.log_metric("test_auc", auc)
    mlflow.sklearn.log_model(best, "model")
""", size=12)
rect(s, 8.4, 3.05, 4.43, 3.75, fill=CARD)
text(s, 8.65, 3.2, 4.0, 0.4, [[("Why track?", {"bold": True, "color": WHITE})]], size=16)
checks(s, 8.65, 3.7, 4.0, 3.0, ["Reproduce any past result", "Compare runs side by side",
                                "Hand the exact model to deployment", "Audit trail for model risk teams"],
       size=13, mark_color=ORANGE, space=8)
notes(s, """
1:45 - Experiment tracking (20 min for this and the next slide, including hands-on).
Stretch idea for fast participants: log every Optuna trial as a nested run using mlflow.start_run(nested=True) inside the objective.
""")

# ---------- MLflow UI
s = base("Comparing Runs in the MLflow UI", "From many experiments to one candidate model", T_M2)
ui = [("1", "Launch", "mlflow ui"), ("2", "Open", "127.0.0.1:5000"), ("3", "Select", "the experiment"),
      ("4", "Tick runs", "then Compare"), ("5", "Sort", "by test_auc"), ("6", "Register", "the best model")]
for i, (n, k, d) in enumerate(ui):
    x0 = 0.5 + i * 2.09
    rect(s, x0, 1.6, 1.9, 1.9, fill=CARD)
    oval(s, x0 + 0.68, 1.75, 0.55, ORANGE, n, color=BG, size=15)
    text(s, x0 + 0.1, 2.4, 1.7, 1.0, [[(k, {"bold": True, "color": WHITE})], [(d, {"color": CYAN, "font": MONO, "size": 11})]],
         size=14, align=PP_ALIGN.CENTER, space_after=2)
    if i < 5:
        arrow(s, x0 + 1.92, 2.55, x0 + 2.07, 2.55, color=DIM)
table(s, 0.5, 3.85, 7.8, [
    ["Run", "max_depth", "n_estimators", "test_auc", "Decision"],
    ["rf_baseline", "None", "100", "0.781", "Reference"],
    ["optuna_trial_17", "11", "240", "0.812", ("Candidate", GREEN)],
    ["optuna_trial_22", "19", "290", "0.806", ("Slower, no gain", ORANGE)],
], [2.1, 1.3, 1.5, 1.2, 1.7], row_h=0.5, size=12)
text(s, 0.5, 5.9, 7.8, 0.35, "Illustrative values: your numbers will differ.", size=10, color=DIM)
rect(s, 8.6, 3.85, 4.23, 2.95, fill=CARD2)
pill(s, 8.85, 4.05, 1.6, "CHECKPOINT", color=GREEN)
text(s, 8.85, 4.55, 3.8, 2.1, [[("main.py trains from the command line, and your best Optuna trial appears in the MLflow UI.",
                                {"bold": True, "color": WHITE})]], size=15)
notes(s, """
Participants open the MLflow UI and find their run. Show the Compare view and the parallel-coordinates plot.
Discuss trial 22: higher complexity with no gain is not a better model - prefer simpler when scores tie (cheaper to serve, easier to explain).
""")

# ---------- M2 checkpoint
s = quiz_checkpoint(T_M2, "main.py trains from the CLI and the best Optuna trial is in MLflow.",
                    ["Ran main.py with --config", "Changed config, reproduced a run",
                     "Optuna study finished", "Best run logged with params, metric, model"],
                    [("Why does random search often beat grid search?",
                      "Few hyperparameters really matter; random search tries more distinct values of each."),
                     ("Why must the test set never guide tuning?",
                      "It becomes part of training, so the final score is optimistic."),
                     ("Name three things every MLflow run should log.",
                      "Parameters, metrics and artefacts (plus tags such as data version).")], ORANGE)
notes(s, "2:00 - Module 2 checkpoint (5 min).")

# ================================================================ MODULE 3
s = divider(3, "Explainability & Drift Monitoring", 55, "2:05 – 3:00", "(e) and (f)",
            [("Explaining models with SHAP", 20), ("From explanation to business narrative", 10),
             ("Detecting data drift: PSI and KS", 15), ("Anomaly detection & alert thresholds", 10)], INDIGO, T_M3)
notes(s, "2:05 - Module 3 starts. Outcomes (e) SHAP explanations in business language and (f) drift detection with PSI and KS.")

# ---------- Why XAI
s = base("Why Explainability Matters", "\"Your model denied my loan application. Why?\"", T_M3)
drivers = [("LAW", "Regulatory compliance", "GDPR Art. 22, MAS FEAT and similar rules expect explanations for automated decisions.", CYAN),
           ("TRUST", "Stakeholder trust", "Underwriters will not adopt a model they cannot question.", GREEN),
           ("DEBUG", "Model debugging", "Catch shortcuts, e.g. a vision model that learned to detect hospital rulers.", ORANGE),
           ("BIAS", "Fairness", "Detect proxies for protected attributes such as age or gender.", RED),
           ("AUDIT", "Auditability", "Keep a reproducible evidence trail for past decisions.", INDIGO)]
for i, (tagx, k, d, col) in enumerate(drivers):
    y0 = 1.55 + i * 0.78
    rect(s, 0.5, y0, 7.4, 0.68, fill=CARD, radius=0.08)
    pill(s, 0.65, y0 + 0.16, 0.95, tagx, color=col, h=0.36, size=10)
    text(s, 1.8, y0, 6.0, 0.68, [[(k + "  ", {"bold": True, "color": WHITE}), (d, {"color": MUTED, "size": 12})]],
         size=13, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 8.2, 1.55, 4.63, 3.8, fill=CARD)
text(s, 8.45, 1.7, 4.2, 0.35, [[("THE TRANSPARENCY SPECTRUM", {"bold": True, "color": DIM})]], size=11)
spec = [("Linear / logistic", "Glass box: read the weights", GREEN), ("Decision trees / GAMs", "Readable rules", CYAN),
        ("Random forest / XGBoost", "Needs post-hoc SHAP", ORANGE), ("Neural nets / LLMs", "Black box: SHAP + LIME", RED)]
for i, (k, d, col) in enumerate(spec):
    y0 = 2.15 + i * 0.78
    rect(s, 8.45, y0, 4.15, 0.68, fill=CARD2, radius=0.08)
    text(s, 8.6, y0, 3.9, 0.68, [[(k, {"bold": True, "color": col})], [(d, {"color": MUTED, "size": 11})]], size=13,
         anchor=MSO_ANCHOR.MIDDLE, space_after=0)
rect(s, 0.5, 5.55, 12.33, 1.25, fill=CARD2, radius=0.08)
text(s, 0.75, 5.55, 11.9, 1.25, [[("Modern answer: ", {"bold": True, "color": CYAN}),
                                  ("train the best-performing model, then explain it post-hoc with SHAP. You no longer have to trade "
                                   "accuracy for compliance.", {"color": WHITE})]], size=15, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
2:05 - Explaining models with SHAP (20 min for this and the next two slides).
Condensed from Session 05 slides 3-4. Keep this to 3 minutes; the hands-on is what matters.
""")

# ---------- SHAP theory + code
s = base("SHAP: Fair Credit for Every Feature", "Shapley values from cooperative game theory (Lundberg & Lee, 2017)", T_M3)
rect(s, 0.5, 1.55, 6.0, 5.25, fill=CARD)
text(s, 0.8, 1.75, 5.4, 0.4, [[("The intuition", {"bold": True, "color": WHITE})]], size=18)
text(s, 0.8, 2.2, 5.4, 1.0, "Features are players in a team; the prediction is the score. SHAP gives each feature its average "
     "marginal contribution across every possible team line-up.", size=14, color=MUTED)
rect(s, 0.8, 3.3, 5.4, 0.8, fill=CODE_BG, radius=0.08, line=LINE)
text(s, 0.8, 3.3, 5.4, 0.8, [[("f(x) = E[f(X)] + Σ φᵢ", {"bold": True, "color": CYAN, "font": MONO})]], size=20,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
checks(s, 0.8, 4.3, 5.5, 2.4, [
    "Additivity: contributions sum exactly to the prediction",
    "Consistency: more influence never means a lower value",
    "Missingness: a feature with no effect gets exactly 0",
], size=14, mark_color=INDIGO, space=10)
code(s, 6.8, 1.55, 6.03, 5.25, """
import shap

explainer = shap.TreeExplainer(model)
shap_values = explainer.shap_values(X_test)

# Global: which features drive the portfolio?
shap.summary_plot(shap_values, X_test)

# Local: why this one applicant?
shap.waterfall_plot(shap.Explanation(
    values=shap_values[0],
    base_values=explainer.expected_value,
    data=X_test.iloc[0]))
""", size=13)
notes(s, """
TreeExplainer is exact and fast for tree ensembles (polynomial time), so it is the default choice for our random forest / gradient boosting models.
Hands-on: participants run the SHAP cells in day1/05_explainability_xai/notebook.ipynb.
""")

# ---------- Global vs local
s = base("Global and Local Explanations", "Same model, two questions: what drives the portfolio, and why this applicant?", T_M3)
for i, (img, k, d, col) in enumerate([
    ("shap_beeswarm.png", "GLOBAL · Beeswarm", "years_employed, age and missed_payments drive portfolio risk. Red = high value; right = more risk.", CYAN),
    ("shap_waterfall.png", "LOCAL · Waterfall", "Applicant #0: baseline 0.26 falls to −2.70 log-odds (≈6.3% default). Age and a clean record pull risk down: approve.", GREEN)]):
    x0 = 0.5 + i * 6.33
    rect(s, x0, 1.55, 6.0, 5.25, fill=CARD)
    pill(s, x0 + 0.25, 1.75, 2.3, k, color=col)
    rect(s, x0 + 0.25, 2.25, 5.5, 2.95, fill=WHITE, radius=0.08)
    pic_w = 5.3
    pic_h = pic_w * (374 / 760 if i == 0 else 453 / 908)
    s.shapes.add_picture(str(ASSETS / img), Inches(x0 + 0.35), Inches(2.25 + (2.95 - pic_h) / 2), Inches(pic_w), Inches(pic_h))
    text(s, x0 + 0.25, 5.35, 5.5, 1.35, d, size=13, color=WHITE)
notes(s, """
Global (beeswarm): every dot is one applicant. Ask: 'What does a red dot far to the right on missed_payments mean?' - many missed payments, strong push towards default.
Local (waterfall): walk the arithmetic. 0.26 + (-0.07 + 0.11 - 0.22 - 0.22 - 0.68 - 1.88) = -2.70. Additivity makes it auditable.
Audience rule: executives get a mean |SHAP| bar chart of the top 5 drivers; underwriters get the waterfall.
""")

# ---------- Business narrative
s = base("From Explanation to Business Narrative", "Never hand raw SHAP numbers to a stakeholder", T_M3)
rect(s, 0.5, 1.5, 12.33, 0.95, fill=CARD)
text(s, 0.75, 1.5, 11.9, 0.95, [[("THE 3-STEP FORMULA", {"bold": True, "color": DIM, "size": 10})],
                                [("Observed value", {"bold": True, "color": WHITE}), ("   →   ", {"color": CYAN}),
                                 ("Direction and size of impact", {"bold": True, "color": WHITE}), ("   →   ", {"color": CYAN}),
                                 ("Actionable next step (counterfactual)", {"bold": True, "color": WHITE})]],
     size=16, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
rect(s, 0.5, 2.65, 7.2, 4.15, fill=CARD)
text(s, 0.75, 2.8, 6.8, 0.4, [[("Worked example: adverse action", {"bold": True, "color": ORANGE})]], size=16)
rect(s, 0.75, 3.3, 6.7, 0.6, fill=CODE_BG, radius=0.08)
text(s, 0.9, 3.3, 6.5, 0.6, [[("credit_utilization = 0.87   SHAP = +0.23", {"color": CYAN, "font": MONO})]], size=13,
     anchor=MSO_ANCHOR.MIDDLE)
rect(s, 0.75, 4.05, 6.7, 1.45, fill=CARD2, radius=0.08)
text(s, 0.9, 4.05, 6.45, 1.45, [[("\"High card utilisation (87%) was the main reason for the decline: it raised this applicant's "
                                  "estimated default risk well above the approval threshold.\"", {"color": GREEN})]],
     size=13, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 0.75, 5.65, 6.7, 0.95, fill=CARD2, radius=0.08)
text(s, 0.9, 5.65, 6.45, 0.95, [[("Counterfactual: ", {"bold": True, "color": CYAN}),
                                 ("\"Reducing utilisation below 30% would move this application to approval.\"", {"color": WHITE})]],
     size=13, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 8.0, 2.65, 4.83, 4.15, fill=CARD)
pill(s, 8.25, 2.85, 2.2, "PAIR EXERCISE · 10 MIN", color=INDIGO)
text(s, 8.25, 3.35, 4.4, 0.9, [[("Write 3 sentences for one applicant from your waterfall plot:", {"bold": True, "color": WHITE})]],
     size=14)
checks(s, 8.25, 4.25, 4.4, 1.6, ["The main driver and its value", "Its direction and size, in plain words",
                                 "One action that would change the outcome"], size=13, mark="›", mark_color=INDIGO)
text(s, 8.25, 5.8, 4.4, 0.9, [[("Swap with a neighbour: ", {"bold": True, "color": ORANGE}),
                               ("could a branch manager act on it with no ML background?", {"color": MUTED})]], size=13)
notes(s, """
2:25 - From explanation to business narrative (10 min). 6 minutes writing, 4 minutes pair review.
Collect one or two strong examples and read them aloud. Checkpoint artefact for Module 3, part 1: one written plain-language explanation.
Avoid: feature names without values ('credit score was important'), jargon ('log-odds'), and no next step.
""")

# ---------- Drift types
s = base("Three Kinds of Drift", "The world moves on after training; your model does not", T_M3)
drift = [("Feature drift", "P(X) changes", "The inputs shift.", "A new digital channel brings much younger applicants.",
          "PSI and KS on every input", CYAN),
         ("Concept drift", "P(y | X) changes", "The rules change.", "After rate rises, the same debt-to-income ratio now means higher risk.",
          "Performance on newly labelled loans", ORANGE),
         ("Label drift", "P(y) changes", "The outcome rate shifts.", "The default rate doubles from 3% to 6% in a downturn.",
          "Predicted vs actual default rate", RED)]
for i, (k, formula, d, ex, detect, col) in enumerate(drift):
    x0 = 0.5 + i * 4.18
    rect(s, x0, 1.55, 3.98, 5.25, fill=CARD)
    text(s, x0 + 0.25, 1.75, 3.5, 0.45, [[(k, {"bold": True, "color": col})]], size=20)
    rect(s, x0 + 0.25, 2.3, 3.48, 0.55, fill=CODE_BG, radius=0.08)
    text(s, x0 + 0.25, 2.3, 3.48, 0.55, [[(formula, {"color": col, "font": MONO})]], size=14, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x0 + 0.25, 3.05, 3.5, 0.4, [[(d, {"bold": True, "color": WHITE})]], size=15)
    text(s, x0 + 0.25, 3.5, 3.5, 1.3, ex, size=13, color=MUTED)
    rect(s, x0 + 0.15, 5.3, 3.68, 1.3, fill=CARD2, radius=0.08)
    text(s, x0 + 0.3, 5.3, 3.4, 1.3, [[("DETECT WITH", {"bold": True, "color": DIM, "size": 10})], [(detect, {"color": WHITE})]],
         size=13, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
notes(s, """
2:35 - Detecting data drift (15 min for this and the next two slides).
Key point: feature drift can be measured immediately; concept drift needs labels, and loan defaults take months to arrive. That delay is why input monitoring matters.
""")

# ---------- PSI & KS with chart (simulated)


def calc_psi(expected, actual, buckets=10):
    bp = np.percentile(expected, np.linspace(0, 100, buckets + 1))
    bp[0], bp[-1] = -np.inf, np.inf
    e = np.clip(np.histogram(expected, bp)[0] / len(expected), 1e-4, None)
    a = np.clip(np.histogram(actual, bp)[0] / len(actual), 1e-4, None)
    return float(np.sum((a - e) * np.log(a / e)))


drng = np.random.default_rng(42)
n_tr, n_pr = 5000, 2000
sim = {
    "credit_utilization": (drng.beta(2, 5, n_tr), drng.beta(2.7, 4.6, n_pr)),
    "age": (drng.normal(40, 10, n_tr), drng.normal(36.5, 10, n_pr)),
    "annual_income": (drng.lognormal(11, 0.4, n_tr), drng.lognormal(10.9, 0.43, n_pr)),
    "loan_amount": (drng.lognormal(9.5, 0.5, n_tr), drng.lognormal(9.52, 0.5, n_pr)),
    "credit_score": (drng.normal(680, 50, n_tr), drng.normal(677, 50, n_pr)),
    "employment_years": (drng.gamma(2, 3, n_tr), drng.gamma(2, 3, n_pr)),
}
drift_rows = []
for f, (e, a) in sim.items():
    ks = ks_2samp(e, a)
    drift_rows.append((f, calc_psi(e, a), float(ks.statistic), float(ks.pvalue)))


def status(p):
    return ("ALERT", RED) if p >= 0.2 else (("WATCH", ORANGE) if p >= 0.1 else ("STABLE", GREEN))


s = base("Measuring Drift: PSI and the KS Test", "Training data vs one simulated production month", T_M3)
rect(s, 0.5, 1.55, 5.6, 5.25, fill=CARD)
text(s, 0.75, 1.7, 5.2, 0.35, [[("POPULATION STABILITY INDEX", {"bold": True, "color": DIM})]], size=11)
rect(s, 0.75, 2.1, 5.1, 0.75, fill=CODE_BG, radius=0.08, line=LINE)
text(s, 0.75, 2.1, 5.1, 0.75, [[("PSI = Σ (Aᵢ − Eᵢ) · ln(Aᵢ / Eᵢ)", {"bold": True, "color": CYAN, "font": MONO})]], size=15,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.75, 2.9, 5.1, 0.6, "Eᵢ, Aᵢ = share of training and production rows in each of 10 training-quantile bins.",
     size=12, color=MUTED)
for i, (lab, rng_, col) in enumerate([("STABLE", "< 0.10", GREEN), ("WATCH", "0.10 – 0.20", ORANGE), ("ALERT", "≥ 0.20", RED)]):
    x0 = 0.75 + i * 1.72
    rect(s, x0, 3.55, 1.6, 0.85, fill=CARD2, radius=0.08)
    text(s, x0, 3.58, 1.6, 0.8, [[(lab, {"bold": True, "color": col})], [(rng_, {"color": WHITE})]], size=13,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=1)
text(s, 0.75, 4.6, 5.1, 0.35, [[("KOLMOGOROV-SMIRNOV TEST", {"bold": True, "color": DIM})]], size=11)
text(s, 0.75, 4.95, 5.1, 1.8, [
    [("D = max | F_train(x) − F_prod(x) |", {"color": CYAN, "font": MONO})],
    [("p < 0.05 means the distributions differ. With thousands of rows KS flags tiny shifts, so pair it with PSI for size.",
      {"color": MUTED})]], size=13, space_after=6)
rect(s, 6.4, 1.55, 6.43, 5.25, fill=CARD)
text(s, 6.65, 1.7, 6.0, 0.35, [[("PSI BY FEATURE (SIMULATED MONTH)", {"bold": True, "color": DIM})]], size=11)
cd = CategoryChartData()
cd.categories = [r[0] for r in drift_rows]
cd.add_series("PSI", [round(r[1], 3) for r in drift_rows])
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(6.55), Inches(2.05), Inches(6.15), Inches(4.6), cd).chart
_dark_chart(ch, 11)
ch.has_legend = False
plot = ch.plots[0]
plot.gap_width = 60
plot.has_data_labels = True
plot.data_labels.number_format = "0.000"
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.color.rgb = rgb(WHITE)
plot.data_labels.font.size = Pt(11)
for j, r in enumerate(drift_rows):
    pt = plot.series[0].points[j]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(status(r[1])[1])
ch.category_axis.reverse_order = True
_axis_dark(ch.category_axis, grid=False)
ch.category_axis.tick_labels.font.size = Pt(11)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0, 0.35
ch.value_axis.major_unit = 0.1
_axis_dark(ch.value_axis)
notes(s, """
The chart is computed by build_deck.py from a seeded simulation (training n=5,000 vs production n=2,000), using the same PSI function as the Session 06 notebook.
credit_utilization crosses the alert line; age is in the watch band. Ask: 'What business event could cause both?' - e.g. a marketing push to younger, more leveraged customers.
""")

# ---------- PSI hands-on table
s = base("Hands-On: Build Your PSI Table", "day1/06_data_drift_anomaly/notebook.ipynb", T_M3)
code(s, 0.5, 1.5, 6.1, 5.3, """
from scipy.stats import ks_2samp

rows = []
for col in numeric_features:
    psi = calculate_psi(train[col], prod[col])
    ks = ks_2samp(train[col], prod[col])
    rows.append({
        "feature": col,
        "psi": round(psi, 3),
        "ks_p": ks.pvalue,
        "status": "ALERT" if psi >= 0.2
                  else "WATCH" if psi >= 0.1
                  else "STABLE",
    })

drift = pd.DataFrame(rows).sort_values("psi", ascending=False)
""", size=12)
trows = [["Feature", "PSI", "KS p-value", "Status"]]
for f, p, d, pv in sorted(drift_rows, key=lambda r: -r[1]):
    lab, col = status(p)
    trows.append([f, f"{p:.3f}", "< 0.001" if pv < 0.001 else f"{pv:.3f}", (lab, col)])
table(s, 6.9, 1.5, 5.93, trows, [2.33, 1.0, 1.4, 1.2], row_h=0.5, size=13)
rect(s, 6.9, 5.2, 5.93, 1.6, fill=CARD2, radius=0.08)
text(s, 7.1, 5.2, 5.6, 1.6, [[("Notice: ", {"bold": True, "color": CYAN}),
                              ("annual_income and credit_score fail the KS test (p < 0.05) but their PSI is small. "
                               "Statistically different is not the same as big enough to act on.", {"color": WHITE})]],
     size=13, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Participants build this table in the Session 06 notebook (simulate a shift, compute PSI and KS per feature).
Checkpoint artefact for Module 3, part 2: a PSI table with drifted features flagged.
""")

# ---------- Isolation Forest & thresholds
s = base("Anomalies, Alert Thresholds and Cadence", "Trainer demonstration: Isolation Forest, then agree a monitoring plan", T_M3)
rect(s, 0.5, 1.55, 5.6, 5.25, fill=CARD)
pill(s, 0.75, 1.75, 2.3, "TRAINER DEMO", color=INDIGO)
text(s, 0.75, 2.25, 5.1, 1.3, [[("Isolation Forest", {"bold": True, "color": WHITE, "size": 18})],
                               [("Random splits isolate unusual rows in very few steps. No labels needed, so it works "
                                 "on live traffic from day one.", {"color": MUTED})]], size=13, space_after=6)
code(s, 0.75, 3.75, 5.1, 2.85, """
from sklearn.ensemble import IsolationForest

iso = IsolationForest(contamination=0.05,
                      random_state=42)
iso.fit(X_train_numeric)

flags = iso.predict(X_new)   # -1 = anomaly
""", size=12)
table(s, 6.4, 1.55, 6.43, [
    ["Cadence", "Check", "Threshold / action"],
    ["Every request", "Schema and range validation", "Reject with HTTP 422"],
    ["Daily", "Volume, missing rate, anomaly rate", "Anomaly rate > 2x baseline"],
    ["Weekly", "PSI and KS per feature", ("Watch 0.1 · Alert 0.2", ORANGE)],
    ["Monthly", "AUC and recall on labelled loans", ("Below SLO: retrain", RED)],
    ["Quarterly", "Model review with risk team", "Re-approve or retire"],
], [1.45, 2.6, 2.38], row_h=0.62, size=12)
rect(s, 6.4, 5.5, 6.43, 1.3, fill=CARD2, radius=0.08)
text(s, 6.6, 5.5, 6.1, 1.3, [[("Alert only when a human must act. ", {"bold": True, "color": CYAN}),
                             ("Everything else goes on the dashboard, not into someone's inbox.", {"color": WHITE})]],
     size=14, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
2:50 - Anomaly detection and alert thresholds (10 min). Run the Isolation Forest cells from the Session 06 notebook live (inject 5% anomalies, show how many are caught).
Then ask the room to agree the cadence table for the loan model. The thresholds shown are starting points; teams tune them to avoid alert fatigue.
""")

# ---------- M3 checkpoint
s = quiz_checkpoint(T_M3, "One plain-language explanation of a prediction, and a PSI table with drifted features flagged.",
                    ["SHAP summary and waterfall plots run", "3-sentence explanation peer-reviewed",
                     "PSI computed for every numeric feature", "ALERT and WATCH features flagged"],
                    [("What does SHAP additivity guarantee?",
                      "Feature contributions add up exactly to the prediction minus the baseline."),
                     ("A feature has PSI = 0.27. What do you do?",
                      "Alert: investigate the cause and consider retraining."),
                     ("Does Isolation Forest need labelled anomalies?",
                      "No. It is unsupervised, so it works on live data immediately.")], INDIGO)
notes(s, "2:55 - Module 3 checkpoint (5 min).")

# ================================================================ MODULE 4
s = divider(4, "Serving, Monitoring & Cloud", 60, "3:00 – 4:00", "(g) and (h)",
            [("Serving a model as a containerised API", 20), ("The MLOps dashboard", 15),
             ("Moving to Azure Machine Learning", 15), ("Action plan and wrap-up", 10)], RED, T_M4)
notes(s, "3:00 - Module 4 starts. Outcomes (g) FastAPI in Docker and (h) production monitoring and the Azure ML workflow.")

# ---------- Serving architecture
s = base("Serving a Model as an API", "The API is the waiter: clients never walk into the kitchen", T_M4)
rect(s, 0.5, 2.3, 2.3, 2.2, fill=CARD)
text(s, 0.5, 2.3, 2.3, 2.2, [[("Client", {"bold": True, "color": WHITE, "size": 17})],
                             [("Web app, mobile app or curl", {"color": MUTED})]], size=13, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 2.85, 3.0, 3.75, 3.0, color=CYAN)
text(s, 2.75, 2.45, 1.1, 0.5, [[("POST JSON", {"color": CYAN, "bold": True})]], size=10, align=PP_ALIGN.CENTER)
arrow(s, 3.75, 3.8, 2.85, 3.8, color=GREEN)
text(s, 2.75, 3.85, 1.1, 0.5, [[("JSON reply", {"color": GREEN, "bold": True})]], size=10, align=PP_ALIGN.CENTER)
rect(s, 3.85, 1.55, 8.98, 3.7, fill=None, radius=0.15, line=CYAN)
text(s, 4.05, 1.65, 4, 0.35, [[("DOCKER CONTAINER", {"bold": True, "color": CYAN})]], size=11)
inner = [("FastAPI", "/health\n/v1/predict\n/docs (Swagger)", WHITE),
         ("Pydantic", "Validates types and ranges; bad input → 422", ORANGE),
         ("Model", "Loaded once at startup, kept in memory", GREEN)]
for i, (k, d, col) in enumerate(inner):
    x0 = 4.1 + i * 2.9
    rect(s, x0, 2.15, 2.6, 2.85, fill=CARD)
    text(s, x0 + 0.15, 2.3, 2.3, 2.6, [[(k, {"bold": True, "color": col, "size": 17})]] +
         [[(ln, {"color": MUTED})] for ln in d.split("\n")], size=13, space_after=4)
    if i < 2:
        arrow(s, x0 + 2.62, 3.55, x0 + 2.88, 3.55, color=DIM)
facts = [("Validate everything", "Never trust input from the internet.", ORANGE),
         ("Load once", "Loading the model per request adds seconds of latency.", GREEN),
         ("Health endpoint", "Load balancers poll /health to route traffic.", CYAN)]
for i, (k, d, col) in enumerate(facts):
    x0 = 0.5 + i * 4.18
    rect(s, x0, 5.5, 3.98, 1.3, fill=CARD2, radius=0.08)
    text(s, x0 + 0.2, 5.5, 3.6, 1.3, [[(k, {"bold": True, "color": col})], [(d, {"color": WHITE})]], size=13,
         anchor=MSO_ANCHOR.MIDDLE, space_after=3)
notes(s, """
3:00 - Serving a model as a containerised API (20 min for this and the next two slides).
Analogy: client = diner, API = waiter, model = chef. Docker ships the whole kitchen so it works identically everywhere.
""")

# ---------- FastAPI code
s = base("FastAPI + Pydantic in 25 Lines", "day1/07_api_serving_containers/app/main.py", T_M4)
code(s, 0.5, 1.5, 7.9, 5.3, """
from fastapi import FastAPI
from pydantic import BaseModel, Field

class PredictRequest(BaseModel):
    age: int = Field(..., ge=18, le=120)
    income: float = Field(..., gt=0)
    credit_score: int = Field(..., ge=300, le=850)

class PredictResponse(BaseModel):
    prediction: str
    probability: float

app = FastAPI(title="Bank Risk ML API")

@app.get("/health")
async def health_check():
    return {"status": "healthy"}

@app.post("/v1/predict", response_model=PredictResponse)
async def predict(req: PredictRequest):
    if req.income > 60000 and req.credit_score > 650:   # stand-in model
        return PredictResponse(prediction="No Default", probability=0.1)
    return PredictResponse(prediction="Default", probability=0.85)
""", size=12)
rect(s, 8.7, 1.5, 4.13, 5.3, fill=CARD)
text(s, 8.95, 1.7, 3.7, 0.4, [[("What to point out", {"bold": True, "color": WHITE})]], size=16)
checks(s, 8.95, 2.2, 3.75, 2.6, [
    "Field(ge=18, le=120) rejects age 7 with HTTP 422, before the model runs",
    "response_model fixes the output contract",
    "/v1/ in the path lets v2 run side by side",
    "Swagger UI appears at /docs for free",
], size=13, mark_color=RED, space=8)
rect(s, 8.95, 5.15, 3.65, 1.45, fill=CARD2, radius=0.08)
text(s, 9.1, 5.15, 3.4, 1.45, [[("Production upgrade: ", {"bold": True, "color": ORANGE}),
                               ("replace the stand-in rule with joblib.load() of your Module 1 Pipeline at startup.",
                                {"color": WHITE})]], size=12, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
The repository app uses a rule-based stand-in model on purpose, so the Docker exercise works without training first.
Demo: uvicorn main:app --reload, open http://localhost:8000/docs, send age=7 to show the 422 validation error.
""")

# ---------- Docker
s = base("Hands-On: Build, Run and Call the Container", "Image = frozen blueprint  ·  Container = running copy", T_M4)
code(s, 0.5, 1.5, 5.6, 4.0, """
FROM python:3.11-slim
WORKDIR /app

# copy requirements first so the
# dependency layer is cached
COPY requirements.txt .
RUN pip install --no-cache-dir -r requirements.txt

COPY . /app
EXPOSE 8000
CMD ["uvicorn", "main:app",
     "--host", "0.0.0.0", "--port", "8000"]
""", size=12, label="Dockerfile")
code(s, 6.4, 1.5, 6.43, 4.0, """
$ cd day1/07_api_serving_containers/app
$ docker build -t loan-api:1.0 .
$ docker run -d -p 8000:8000 --name loan-api loan-api:1.0

$ curl -X POST http://localhost:8000/v1/predict \\
    -H "Content-Type: application/json" \\
    -d '{"age": 35, "income": 72000, "credit_score": 710}'
→ {"prediction":"No Default","probability":0.1}
""", size=12, label="TERMINAL")
rect(s, 0.5, 5.75, 12.33, 1.05, fill=CARD2, radius=0.08)
pill(s, 0.75, 6.1, 1.6, "CHECKPOINT", color=GREEN)
text(s, 2.55, 5.75, 10.1, 1.05, [[("The container returns a prediction to your curl request. ", {"bold": True, "color": WHITE}),
                                  ("Windows PowerShell: use curl.exe and escape the inner quotes, or test from /docs.",
                                   {"color": MUTED})]], size=13, anchor=MSO_ANCHOR.MIDDLE)
notes(s, """
Participants build and run the container (about 8 minutes; the first build downloads the python:3.11-slim base image, so ask people to pre-pull it at the start of the session if Wi-Fi is slow).
Troubleshooting: port 8000 already in use -> docker run -p 8001:8000; Docker Desktop not running -> start it and retry.
Cleanup after: docker rm -f loan-api.
""")

# ---------- Dashboard (6 panels with sparklines)
s = base("The Six-Panel MLOps Dashboard", "Four pillars on one screen: model, data, system and business", T_M4)


def spark(slide, x, y, w, h, values, col, kind="line", threshold=None):
    cd_ = CategoryChartData()
    cd_.categories = [str(i) for i in range(len(values))]
    cd_.add_series("v", values)
    ct = XL_CHART_TYPE.LINE if kind == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
    if threshold is not None:
        cd_.add_series("t", [threshold] * len(values))
    c = slide.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd_).chart
    _dark_chart(c, 9)
    c.has_legend = False
    c.category_axis.visible = False
    _axis_dark(c.value_axis)
    c.value_axis.tick_labels.font.size = Pt(9)
    p_ = c.plots[0]
    s0 = p_.series[0]
    if kind == "line":
        s0.format.line.color.rgb = rgb(col)
        s0.format.line.width = Pt(2.25)
        s0.smooth = False
        s0.marker.style = XL_MARKER_STYLE.NONE
        if threshold is not None:
            s1 = p_.series[1]
            s1.format.line.color.rgb = rgb(RED)
            s1.format.line.width = Pt(1.5)
            s1.format.line.dash_style = 4
            s1.marker.style = XL_MARKER_STYLE.NONE
    else:
        p_.gap_width = 40
        s0.format.fill.solid()
        s0.format.fill.fore_color.rgb = rgb(col)
    return c


panels6 = [
    ("MODEL", "AUC over time vs SLO", CYAN, "line", [0.82, 0.82, 0.81, 0.81, 0.80, 0.79, 0.78, 0.76, 0.75, 0.73], 0.75, "Alert below 0.75"),
    ("MODEL", "Predicted default rate", CYAN, "line", [3.1, 3.0, 3.2, 3.3, 3.6, 3.9, 4.2, 4.6, 4.9, 5.3], None, "Compare with actuals"),
    ("DATA", "Top-feature PSI this week", INDIGO, "bar", [0.30, 0.14, 0.06, 0.01, 0.01, 0.01], None, "Alert at 0.2"),
    ("SYSTEM", "p99 latency (ms)", GREEN, "line", [110, 120, 115, 130, 125, 180, 240, 210, 150, 140], 300, "SLO 300 ms"),
    ("SYSTEM", "Requests & error rate", GREEN, "bar", [820, 900, 870, 950, 1010, 990, 1100, 1050, 980, 1020], None, "Errors < 1%"),
    ("BUSINESS", "Approved loans / expected loss", ORANGE, "line", [100, 102, 101, 104, 103, 99, 96, 94, 92, 90], None, "Watch the trend"),
]
for i, (pillar, k, col, kind, vals, thr, foot) in enumerate(panels6):
    r_, c_ = divmod(i, 3)
    x0, y0 = 0.5 + c_ * 4.18, 1.5 + r_ * 2.68
    rect(s, x0, y0, 3.98, 2.5, fill=CARD)
    pill(s, x0 + 0.2, y0 + 0.15, 1.15, pillar, color=col, h=0.3, size=9)
    text(s, x0 + 1.45, y0 + 0.12, 2.45, 0.38, [[(k, {"bold": True, "color": WHITE})]], size=12, anchor=MSO_ANCHOR.MIDDLE)
    spark(s, x0 + 0.1, y0 + 0.5, 3.78, 1.6, vals, col, kind, thr)
    text(s, x0 + 0.2, y0 + 2.08, 3.6, 0.35, foot, size=10, color=DIM)
notes(s, """
3:20 - The MLOps dashboard (15 min for this and the next slide). Values are illustrative.
Read the dashboard as a story: AUC is sliding towards the SLO line, predicted default rate is rising, and credit_utilization PSI is in alert - consistent with the drift found in Module 3. Latency spiked mid-period but stayed within SLO.
Hands-on option: run the six-panel cell in day1/08_mlops_dashboard/notebook.ipynb.
""")

# ---------- Golden signals & SLOs
s = base("Golden Signals, SLOs and Alert Fatigue", "Measure what matters; alert only on broken promises", T_M4)
gs = [("Latency", "How fast? Track p99, not the average.", CYAN), ("Traffic", "How many requests per second?", GREEN),
      ("Errors", "What share of requests fail?", RED), ("Saturation", "How close to full is CPU or memory?", ORANGE)]
for i, (k, d, col) in enumerate(gs):
    x0 = 0.5 + i * 3.13
    rect(s, x0, 1.55, 2.95, 1.6, fill=CARD)
    text(s, x0 + 0.2, 1.7, 2.6, 1.35, [[(k, {"bold": True, "color": col, "size": 18})], [(d, {"color": WHITE})]], size=13,
         space_after=4)
text(s, 0.5, 3.35, 6, 0.35, [[("GOOGLE SRE GOLDEN SIGNALS (ABOVE)  ·  EXAMPLE SLOs FOR THE LOAN MODEL (BELOW)",
                               {"bold": True, "color": DIM})]], size=11)
table(s, 0.5, 3.75, 7.6, [
    ["Service level objective", "Pillar"],
    ["AUC ≥ 0.75 on the monthly labelled sample", "Model"],
    ["PSI < 0.2 on the top-10 features", "Data"],
    ["p99 latency < 300 ms; error rate < 1%", "System"],
    ["Expected loss within 10% of plan", "Business"],
], [5.8, 1.8], row_h=0.6, size=13)
rect(s, 8.4, 3.75, 4.43, 3.05, fill=CARD2)
text(s, 8.65, 3.9, 4.0, 2.8, [[("Alert fatigue", {"bold": True, "color": RED, "size": 18})],
                              [("If every 0.1% dip sends an email, people stop reading them. Alert on SLO breaches only; "
                                "everything else stays on the dashboard.", {"color": WHITE})]], size=14, space_after=8)
notes(s, "Ask: 'Which of these SLOs can you measure on day one, and which need to wait for labels?' (System and data immediately; model performance waits for loan outcomes.)")

# ---------- Azure ML
s = base("Moving the Same Workflow to Azure ML", "Trainer demonstration: no Azure account needed today", T_M4)
table(s, 0.5, 1.5, 7.4, [
    ["What you built today", "Azure Machine Learning equivalent"],
    ["main.py + config.yaml", "Command job on a compute cluster"],
    ["Optuna study", "Sweep job (Bayesian, early termination)"],
    ["Local MLflow tracking", "Built-in MLflow tracking in the workspace"],
    ["model.joblib in artifacts/", "Model registry (versioned, tagged)"],
    ["FastAPI + Docker", "Managed online endpoint"],
    ["PSI / KS notebook checks", "Model monitoring (data drift signals)"],
], [3.1, 4.3], row_h=0.6, size=13, col_colors=[WHITE, CYAN])
code(s, 0.5, 5.85, 7.4, 0.95, """
job = command(code="./scripts", command="python train.py",
    environment="azureml:credit-model-env:1", compute="cpu-cluster")
""", size=11, label=None)
rect(s, 8.2, 1.5, 4.63, 5.3, fill=CARD)
pill(s, 8.45, 1.7, 2.1, "DEMO SCRIPT · 15 MIN", color=RED)
for i, (k, d) in enumerate([("Workspace tour", "Studio, compute, data assets"), ("Submit a command job", "Same training script, cloud compute"),
                            ("Watch metrics live", "MLflow run in the Jobs view"), ("Register the model", "Version, tags, lineage")]):
    y0 = 2.25 + i * 1.1
    oval(s, 8.45, y0 + 0.12, 0.5, RED, str(i + 1), color=BG, size=14)
    text(s, 9.1, y0, 3.6, 0.95, [[(k, {"bold": True, "color": WHITE})], [(d, {"color": MUTED, "size": 12})]], size=14,
         space_after=1)
notes(s, """
3:35 - Moving to Azure Machine Learning (15 min, trainer demonstration). Uses the trainer's workspace and the Day 2 Session 10 assets (environment credit-model-env:1, compute cpu-cluster).
Key message: nothing learned today is thrown away. The script, the MLflow calls and the model artefact move to the cloud unchanged; Azure adds managed compute, a registry and endpoints.
Fallback if the network fails: walk through the screenshots in day2/02_cloud_training_experiment.
""")

# ---------- Action plan
s = base("Your Back-to-Work Action Plan", "Five minutes, individually: commit to specific next steps", T_M4)
plan = [("1", "One pipeline I will audit for leakage", GREEN), ("2", "One script I will refactor into config + modules", ORANGE),
        ("3", "One prediction I will explain to a stakeholder", INDIGO), ("4", "One drift or alert check I will add", INDIGO),
        ("5", "My first step on Monday, and my deadline", RED)]
for i, (n, k, col) in enumerate(plan):
    y0 = 1.55 + i * 0.98
    rect(s, 0.5, y0, 7.6, 0.85, fill=CARD, radius=0.08)
    oval(s, 0.68, y0 + 0.17, 0.5, col, n, color=BG, size=14)
    text(s, 1.35, y0, 4.0, 0.85, [[(k, {"color": WHITE, "bold": True})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 5.45, y0 + 0.2, 2.45, 0.45, fill=CARD2, radius=0.06)
rect(s, 8.4, 1.55, 4.43, 5.25, fill=CARD)
text(s, 8.65, 1.7, 4.0, 0.4, [[("A 30-60-90 day path", {"bold": True, "color": WHITE})]], size=17)
for i, (d, t, col) in enumerate([("30 days", "Leakage audit + Pipeline on one live model", GREEN),
                                 ("60 days", "Scripts, Optuna and MLflow for that model", ORANGE),
                                 ("90 days", "API, drift checks and a dashboard in place", RED)]):
    y0 = 2.3 + i * 1.45
    rect(s, 8.65, y0, 3.95, 1.25, fill=CARD2, radius=0.08)
    text(s, 8.85, y0, 3.6, 1.25, [[(d, {"bold": True, "color": col})], [(t, {"color": WHITE})]], size=14,
         anchor=MSO_ANCHOR.MIDDLE, space_after=2)
notes(s, """
3:50 - Back-to-work action plan and wrap-up (10 min: 5 min writing, 5 min takeaways, questions and feedback).
Hand out the personal action plan template. Ask two volunteers to share item 5.
Checkpoint: each participant leaves with a personal action plan.
""")

# ---------- Key takeaways
s = base("Key Takeaways", "One sentence per module to remember on Monday", T_M4)
tk = [("MODULE 1", "Trustworthy data", ["Pick the tool from the data shape", "Split first; preprocess inside a Pipeline",
                                        "Validate the way the model will be used"], GREEN),
      ("MODULE 2", "Training, tuning, tracking", ["Config + seeds = reproducible runs", "Tune on validation folds with Optuna",
                                                  "If it is not in MLflow, it did not happen"], ORANGE),
      ("MODULE 3", "Explain and monitor", ["SHAP for audits, plain words for people", "PSI ≥ 0.2 means investigate",
                                           "Alert only when a human must act"], INDIGO),
      ("MODULE 4", "Serve and scale", ["Validate input; load the model once", "Dashboards track 4 pillars against SLOs",
                                       "The same workflow runs on Azure ML"], RED)]
for i, (m, k, items, col) in enumerate(tk):
    x0 = 0.5 + i * 3.13
    rect(s, x0, 1.55, 2.95, 4.35, fill=CARD)
    pill(s, x0 + 0.2, 1.75, 1.35, m, color=col, h=0.32, size=10)
    text(s, x0 + 0.2, 2.2, 2.6, 0.75, [[(k, {"bold": True, "color": WHITE})]], size=19)
    checks(s, x0 + 0.2, 2.95, 2.6, 2.9, items, size=15, mark_color=col, space=14)
rect(s, 0.5, 6.1, 12.33, 0.7, fill=CARD2, radius=0.08)
text(s, 0.75, 6.1, 11.9, 0.7, [[("Questions?   ", {"bold": True, "color": CYAN}),
                                ("Please complete the programme feedback form before you leave.", {"color": WHITE})]],
     size=15, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Read one takeaway per module, take final questions, then share the feedback form link.")

# ---------- Closing
s = base()
pill(s, 5.42, 2.2, 2.5, "END OF DAY 1", color=GREEN)
text(s, 1.0, 2.75, 11.33, 1.2, [[("Thank You", {"bold": True, "color": WHITE})]], size=54, align=PP_ALIGN.CENTER)
text(s, 1.0, 3.95, 11.33, 0.5, "From a leaky notebook to a monitored, containerised model in four hours.", size=17,
     align=PP_ALIGN.CENTER)
rect(s, 3.67, 4.75, 6.0, 0.55, fill=CARD, radius=0.08)
text(s, 3.67, 4.75, 6.0, 0.55, [[("Next: Day 2  ·  Azure ML workspace, cloud training and the model registry", {"bold": True, "color": CYAN})]],
     size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "4:00 - Close. Thank participants; remind them that Day 2 builds on today's scripts in Azure ML.")

prs.save(OUT)
print(f"Saved {OUT.name} with {_slide_no[0]} slides")
